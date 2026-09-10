import asyncio
import os
import sqlite3
import threading
import time
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
import httpx
import pytest
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook as ExcelWorkbook
from openpyxl import load_workbook


TEST_DATABASE = Path("/tmp/siamlink-test.db")
os.environ["AI_PROVIDER"] = "demo"
os.environ["APP_DATABASE_PATH"] = str(TEST_DATABASE)

from fastapi.testclient import TestClient

from app import main as main_module
from app.config import settings
from app.database import Database
from app.main import app
from app.services.exports import (
    ExportedDocument,
    _rect_overlap_area,
    _scan_ink_scale,
    build_layout_pdf_export,
    build_pdf_export,
)
from app.services.documents import (
    DocumentSegment,
    ParsedDocument,
    ParsedPdfPage,
    RenderedPdfPage,
    parse_document,
    render_pdf_pages,
    render_pdf_tiles,
)
from app.services.translator import (
    OpenAIProvider,
    SegmentTranslationResult,
    TranslationService,
    TranslationResult,
    detect_language,
    normalize_translation_text,
    parse_indexed_image_sources,
    quality_checks,
    structured_translation_instructions,
    translation_instructions,
)
from app.services.native_pdf import _find_system_font


def setup_module():
    if TEST_DATABASE.exists():
        TEST_DATABASE.unlink()


def test_cad_local_ocr_coordinator_overlaps_locator_with_bounded_paddle():
    async def scenario():
        coordinator = main_module._CadLocalOcrCoordinator(1)
        events = []
        first_paddle_acquired = asyncio.Event()
        release_first_paddle = asyncio.Event()
        locator_acquired = asyncio.Event()
        release_exclusive = asyncio.Event()
        both_paddles_acquired = asyncio.Event()

        async def first_paddle():
            async with coordinator.paddle():
                events.append("paddle-1")
                first_paddle_acquired.set()
                await release_first_paddle.wait()

        async def exclusive_locator():
            async with coordinator.exclusive():
                events.append("locator")
                locator_acquired.set()
                await release_exclusive.wait()

        async def second_paddle():
            async with coordinator.paddle():
                events.append("paddle-2")

        first_task = asyncio.create_task(first_paddle())
        await first_paddle_acquired.wait()
        locator_task = asyncio.create_task(exclusive_locator())
        await asyncio.wait_for(locator_acquired.wait(), timeout=1)
        second_task = asyncio.create_task(second_paddle())
        await asyncio.sleep(0)
        assert events == ["paddle-1", "locator"]

        release_first_paddle.set()
        while events == ["paddle-1", "locator"]:
            await asyncio.sleep(0)
        assert events == ["paddle-1", "locator", "paddle-2"]

        release_exclusive.set()
        await asyncio.gather(first_task, locator_task, second_task)
        assert events == ["paddle-1", "locator", "paddle-2"]

        concurrent = main_module._CadLocalOcrCoordinator(2)
        release_paddles = asyncio.Event()
        active = 0

        async def shared_paddle():
            nonlocal active
            async with concurrent.paddle():
                active += 1
                if active == 2:
                    both_paddles_acquired.set()
                await release_paddles.wait()

        shared_tasks = [asyncio.create_task(shared_paddle()) for _ in range(2)]
        await asyncio.wait_for(both_paddles_acquired.wait(), timeout=1)
        release_paddles.set()
        await asyncio.gather(*shared_tasks)

    asyncio.run(scenario())


def test_cad_indexed_no_text_retry_keeps_short_confident_thai_label():
    candidate = {
        "source_hint": "จุ",
        "source_confidence": 91.0,
    }

    assert main_module._cad_indexed_read_needs_review(
        candidate, {"source_text": "[NO_TEXT]"}
    )
    assert not main_module._cad_indexed_read_needs_review(
        candidate, {"source_text": "จุ"}
    )
    candidate["source_confidence"] = 55.0
    assert not main_module._cad_indexed_read_needs_review(
        candidate, {"source_text": "[NO_TEXT]"}
    )


def test_slow_structured_translation_uses_delayed_hedge(monkeypatch):
    calls = 0

    async def translate_segments(
        segments,
        source_language,
        _target_language,
        _context,
        *,
        require_complete,
    ):
        nonlocal calls
        calls += 1
        call_number = calls
        await asyncio.sleep(0.10 if call_number == 1 else 0.001)
        return SegmentTranslationResult(
            source_language=source_language,
            translations={segment_id: "译文" for segment_id in segments},
            provider=f"provider-{call_number}",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "TEXT_MODEL_HEDGE_DELAY_SECONDS", 0.01)
    monkeypatch.setattr(
        main_module.database,
        "find_matching_knowledge",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", translate_segments
    )
    batch = [
        DocumentSegment(
            segment_id="cad:p1:source:001",
            page_number=1,
            text="ข้อความ",
            source_kind="outline-text",
            bbox=(0, 0, 10, 10),
            metadata={},
        )
    ]

    async def run_batch():
        return await main_module._translate_document_segment_batch(
            batch,
            "th",
            "zh",
            "",
            asyncio.Semaphore(2),
        )

    translations, providers, _warnings = asyncio.run(run_batch())

    assert translations == {"cad:p1:source:001": "译文"}
    assert providers == ["provider-2"]
    assert calls == 2


def test_fast_structured_translation_does_not_duplicate(monkeypatch):
    calls = 0

    async def translate_segments(
        segments,
        source_language,
        _target_language,
        _context,
        *,
        require_complete,
    ):
        nonlocal calls
        calls += 1
        return SegmentTranslationResult(
            source_language=source_language,
            translations={segment_id: "译文" for segment_id in segments},
            provider="provider",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "TEXT_MODEL_HEDGE_DELAY_SECONDS", 0.05)
    monkeypatch.setattr(
        main_module.database,
        "find_matching_knowledge",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", translate_segments
    )
    batch = [
        DocumentSegment(
            segment_id="pdf:p1:u1",
            page_number=1,
            text="source",
            source_kind="paragraph",
            bbox=(0, 0, 10, 10),
            metadata={},
        )
    ]

    async def run_batch():
        await main_module._translate_document_segment_batch(
            batch,
            "en",
            "zh",
            "",
            asyncio.Semaphore(2),
        )

    asyncio.run(run_batch())

    assert calls == 1


def test_partial_slow_translation_does_not_cancel_complete_hedge(monkeypatch):
    calls = 0

    async def translate_segments(
        segments,
        source_language,
        _target_language,
        _context,
        *,
        require_complete,
    ):
        nonlocal calls
        calls += 1
        call_number = calls
        await asyncio.sleep(0.02)
        segment_ids = list(segments)
        returned_ids = segment_ids[:1] if call_number == 1 else segment_ids
        return SegmentTranslationResult(
            source_language=source_language,
            translations={segment_id: "译文" for segment_id in returned_ids},
            provider=f"provider-{call_number}",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "TEXT_MODEL_HEDGE_DELAY_SECONDS", 0.01)
    monkeypatch.setattr(
        main_module.database,
        "find_matching_knowledge",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", translate_segments
    )
    batch = [
        DocumentSegment(
            segment_id=f"pdf:p1:u{index}",
            page_number=1,
            text=f"source {index}",
            source_kind="paragraph",
            bbox=(0, 0, 10, 10),
            metadata={},
        )
        for index in (1, 2)
    ]

    async def run_batch():
        return await main_module._translate_document_segment_batch(
            batch,
            "en",
            "zh",
            "",
            asyncio.Semaphore(2),
        )

    translations, providers, _warnings = asyncio.run(run_batch())

    assert set(translations) == {"pdf:p1:u1", "pdf:p1:u2"}
    assert providers == ["provider-2"]
    assert calls == 2


def make_pdf(page_texts, *, image_only=False):
    document = fitz.open()
    try:
        for text in page_texts:
            page = document.new_page(width=595, height=842)
            if image_only:
                source = fitz.open()
                try:
                    source_page = source.new_page(width=595, height=842)
                    source_page.insert_text((72, 120), text, fontsize=24)
                    pixmap = source_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
                finally:
                    source.close()
            else:
                page.insert_text((72, 120), text, fontsize=14)
        return document.tobytes()
    finally:
        document.close()


def make_native_thai_pdf(page_texts):
    document = fitz.open()
    font_path = str(_find_system_font("zh"))
    try:
        for text in page_texts:
            page = document.new_page(width=595, height=842)
            page.insert_font(fontname="ThaiSource", fontfile=font_path)
            page.insert_text(
                (72, 120),
                text,
                fontname="ThaiSource",
                fontfile=font_path,
                fontsize=14,
            )
        return document.tobytes()
    finally:
        document.close()


def make_mixed_pdf():
    document = fitz.open()
    try:
        text_page = document.new_page(width=595, height=842)
        text_page.insert_text(
            (72, 120),
            "This is a selectable text page with enough content for translation.",
            fontsize=14,
        )

        image_page = document.new_page(width=595, height=842)
        source = fitz.open()
        try:
            source_page = source.new_page(width=595, height=842)
            source_page.insert_text((72, 120), "Image page content", fontsize=24)
            pixmap = source_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_page.insert_image(image_page.rect, stream=pixmap.tobytes("png"))
        finally:
            source.close()
        return document.tobytes()
    finally:
        document.close()


def make_minimal_pptx():
    slide_xml = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree><p:sp><p:txBody>
    <a:bodyPr/><a:lstStyle/>
    <a:p><a:r><a:rPr lang="en-US" sz="2400"/><a:t>Hello slide</a:t></a:r></a:p>
  </p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>'''
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
    return output.getvalue()


KNOWLEDGE_XLSX_HEADERS = [
    "NO",
    "Original Text（Thai）",
    "Translated Text（English）",
    "Revised Text（English）",
    "Translated Text（Chinese）",
    "Revised Text（Chinese）",
]


def make_knowledge_xlsx(rows, *, headers=None, header_row=1):
    workbook = ExcelWorkbook()
    worksheet = workbook.active
    worksheet.title = "修订"
    for _ in range(header_row - 1):
        worksheet.append(["说明"])
    worksheet.append(headers or KNOWLEDGE_XLSX_HEADERS)
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_health_reports_demo_provider():
    with TestClient(app) as client:
        response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["provider"] == "demo"


def test_indexed_image_source_parser_keeps_stable_ids_without_translations():
    assert parse_indexed_image_sources(
        '{"items":[{"id":"ID001","source_text":"ภาษาไทย"}]}'
    ) == [{"id": "ID001", "source_text": "ภาษาไทย"}]


def test_indexed_source_reader_can_return_confirmed_rows_for_compact_retry():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.read_indexed_image_lines.return_value = (
        [{"id": "ID001", "source_text": "ผนัง"}],
        "vision",
    )

    partial, route = asyncio.run(
        service.read_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001", "ID002"],
            require_complete=False,
        )
    )

    assert route == "vision"
    assert partial == [{"id": "ID001", "source_text": "ผนัง"}]
    with pytest.raises(RuntimeError, match="ID_MISMATCH"):
        asyncio.run(
            service.read_indexed_image_lines(
                b"image", "image/png", ["ID001", "ID002"]
            )
        )


def test_grouped_indexed_source_reader_validates_global_ids():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.read_indexed_image_line_group.return_value = (
        [
            {"id": "ID0001", "source_text": "ผนัง"},
            {"id": "ID0002", "source_text": "ประตู"},
        ],
        "vision",
    )

    items, route = asyncio.run(
        service.read_indexed_image_line_group(
            [b"sheet-1", b"sheet-2"],
            "image/png",
            ["ID0001", "ID0002"],
        )
    )

    assert route == "vision"
    assert [item["id"] for item in items] == ["ID0001", "ID0002"]


def test_text_translation_accepts_request_context_and_saves_history():
    with TestClient(app) as client:
        response = client.post(
            "/api/translate",
            json={
                "text": "你好",
                "source_language": "auto",
                "target_language": "th",
                "context": "SiamLink 是产品名",
            },
        )
        history = client.get("/api/history").json()

    assert response.status_code == 200
    assert response.json()["translated_text"] == "สวัสดี"
    assert response.json()["source_language"] == "zh"
    assert response.json()["kind"] == "text"
    assert history[0]["id"] == response.json()["id"]
    assert "context" not in history[0]


def test_image_translation_and_history():
    with TestClient(app) as client:
        response = client.post(
            "/api/translate/image",
            data={"source_language": "th", "target_language": "zh", "context": "菜单"},
            files={"file": ("menu.png", b"not-decoded-in-demo", "image/png")},
        )

    assert response.status_code == 200
    assert response.json()["kind"] == "image"
    assert response.json()["filename"] == "menu.png"
    assert response.json()["source_text"]


def test_document_translation_extracts_text():
    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh", "context": "合同"},
            files={"file": ("agreement.txt", "Hello\nThank you".encode(), "text/plain")},
        )

    assert response.status_code == 200
    assert response.json()["kind"] == "document"
    assert response.json()["source_text"] == "Hello\nThank you"
    assert response.json()["export_filename"] == "agreement-译文.txt"

    with TestClient(app) as client:
        exported = client.get(f"/api/history/{response.json()['id']}/export")

    assert exported.status_code == 200
    assert "演示译文" in exported.text


def test_plain_text_export_is_available_for_every_history_item():
    with TestClient(app) as client:
        response = client.post(
            "/api/translate",
            json={
                "text": "Hello",
                "source_language": "en",
                "target_language": "zh",
            },
        )
        exported = client.get(
            f"/api/history/{response.json()['id']}/export/text"
        )

    assert response.status_code == 200
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("text/plain")
    assert exported.headers["cache-control"] == "private, max-age=86400"
    assert exported.content.decode("utf-8-sig").strip() == "你好"
    assert "translation.txt" in exported.headers["content-disposition"]


def test_pdf_export_keeps_page_count_size_and_page_correspondence():
    content = make_native_thai_pdf(["สวัสดี", "ขอบคุณ"])

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("sample.pdf", content, "application/pdf")},
        )
        exported = client.get(
            f"/api/history/{response.json()['id']}/export"
        )
        inline_export = client.get(
            f"/api/history/{response.json()['id']}/export?inline=true",
            headers={"Range": "bytes=0-99"},
        )

    body = response.json()
    assert response.status_code == 200
    assert body["export_filename"] == "sample-译文.pdf"
    assert exported.status_code == 200
    assert len(exported.content) < 20_000_000
    assert exported.headers["content-disposition"].startswith(
        'attachment; filename="translation.pdf"; filename*=UTF-8\'\''
    )
    assert inline_export.status_code == 206
    assert len(inline_export.content) == 100
    assert inline_export.headers["content-disposition"].startswith(
        'inline; filename="translation.pdf"; filename*=UTF-8\'\''
    )
    assert inline_export.headers["content-range"].startswith("bytes 0-99/")
    source_pdf = fitz.open(stream=content, filetype="pdf")
    translated_pdf = fitz.open(stream=exported.content, filetype="pdf")
    try:
        assert translated_pdf.page_count == source_pdf.page_count == 2
        for source_page, translated_page in zip(source_pdf, translated_pdf):
            assert translated_page.rect.width == source_page.rect.width
            assert translated_page.rect.height == source_page.rect.height
        first_text = translated_pdf[0].get_text()
        second_text = translated_pdf[1].get_text()
        assert "你好" in first_text
        assert "谢谢" not in first_text
        assert "谢谢" in second_text
        assert "你好" not in second_text
        first_blocks = translated_pdf[0].get_text("blocks")
        assert first_blocks
        assert abs(first_blocks[0][0] - 72) < 4
        assert abs(first_blocks[0][1] - 105) < 8
    finally:
        translated_pdf.close()
        source_pdf.close()


def test_layout_pdf_export_supports_original_tiny_font_sizes():
    content = make_pdf(["source"])
    exported = build_layout_pdf_export(
        content,
        [
            {
                "segment_id": "pdf:p1:b1",
                "page_number": 1,
                "text": "R",
                "translated_text": "R",
                "source_kind": "text",
                "bbox": [72, 100, 80, 103],
                "font_size": 2.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {"leading": 2.2},
            }
        ],
        "zh",
    )

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        assert document.page_count == 1
    finally:
        document.close()


def test_scanned_pdf_complete_line_shrinks_instead_of_wrapping():
    content = make_pdf(["source"])
    exported = build_layout_pdf_export(
        content,
        [
            {
                "segment_id": "pdf:p1:o1",
                "page_number": 1,
                "text": "ลงชื่อ",
                "translated_text": "签名",
                "source_kind": "ocr",
                "bbox": [40, 60, 51, 80],
                "font_size": 12.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {
                    "page_type": "image",
                    "visual_pass": True,
                    "translation_unit": "complete-line",
                    "rotation": 0,
                    "leading": 14.0,
                },
            }
        ],
        "zh",
    )

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        matching_lines = [
            "".join(span["text"] for span in line["spans"])
            for block in document[0].get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block.get("lines", [])
            if any(
                character in "".join(span["text"] for span in line["spans"])
                for character in "签名"
            )
        ]
        assert matching_lines == ["签名"]
    finally:
        document.close()


def test_layout_result_reuses_translation_for_identical_ocr_source():
    document = ParsedDocument(
        filename="scan.pdf",
        text="",
        page_count=2,
        warnings=[],
    )
    first = {
        "segment_id": "pdf:p1:o1",
        "page_number": 1,
        "text": "กรรมการกำหนดราคากลาง",
        "translated_text": "基准价制定委员会委员",
        "source_kind": "ocr",
        "bbox": [10, 10, 100, 20],
        "metadata": {"visual_pass": True},
    }
    second = {
        **first,
        "segment_id": "pdf:p2:o1",
        "page_number": 2,
        "translated_text": "参考价格核定委员",
    }

    _, result = main_module._build_layout_translation_result(
        document,
        "th",
        {},
        ["test"],
        [],
        [
            (1, first["text"], None, [first]),
            (2, second["text"], None, [second]),
        ],
    )

    assert [
        item["translated_text"] for item in result.layout_segments
    ] == ["基准价制定委员会委员", "基准价制定委员会委员"]


def test_layout_pdf_export_matches_rotated_source_page_orientation():
    source = fitz.open()
    try:
        page = source.new_page(width=595, height=842)
        page.insert_text((72, 120), "source", fontsize=12)
        page.set_rotation(180)
        content = source.tobytes()
    finally:
        source.close()

    parsed_source = fitz.open(stream=content, filetype="pdf")
    try:
        source_bbox = next(
            span["bbox"]
            for block in parsed_source[0].get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block.get("lines", [])
            for span in line.get("spans", [])
            if span["text"] == "source"
        )
        expected_bbox = fitz.Rect(source_bbox) * parsed_source[0].rotation_matrix
    finally:
        parsed_source.close()

    exported = build_layout_pdf_export(
        content,
        [
            {
                "segment_id": "pdf:p1:b1",
                "page_number": 1,
                "text": "source",
                "translated_text": "UPRIGHT",
                "source_kind": "text",
                "bbox": source_bbox,
                "font_size": 12.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {"leading": 14.0},
            }
        ],
        "en",
    )

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        lines = [
            line
            for block in document[0].get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block.get("lines", [])
            if "UPRIGHT" in "".join(span["text"] for span in line["spans"])
        ]
        assert document[0].rotation == 0
        assert lines
        assert lines[0]["dir"][0] > 0.99
        translated_bbox = fitz.Rect(lines[0]["bbox"])
        assert abs(translated_bbox.x0 - expected_bbox.x0) < 4
        assert abs(translated_bbox.y0 - expected_bbox.y0) < 8
        assert "source" not in document[0].get_text("text")
    finally:
        document.close()


def test_layout_pdf_export_accepts_legacy_display_coordinates():
    source = fitz.open()
    try:
        page = source.new_page(width=595, height=842)
        page.insert_text((72, 120), "source", fontsize=12)
        page.set_rotation(180)
        content = source.tobytes()
    finally:
        source.close()

    parsed_source = fitz.open(stream=content, filetype="pdf")
    try:
        source_bbox = next(
            span["bbox"]
            for block in parsed_source[0].get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block.get("lines", [])
            for span in line.get("spans", [])
            if span["text"] == "source"
        )
        display_bbox = list(fitz.Rect(source_bbox) * parsed_source[0].rotation_matrix)
    finally:
        parsed_source.close()

    exported = build_layout_pdf_export(
        content,
        [
            {
                "segment_id": "cad:p1:l1",
                "page_number": 1,
                "text": "source",
                "translated_text": "LEGACY",
                "source_kind": "text",
                "bbox": display_bbox,
                "font_size": 12.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {"leading": 14.0},
            }
        ],
        "en",
    )

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        assert document[0].rotation == 0
        assert "LEGACY" in document[0].get_text()
    finally:
        document.close()


def test_layout_pdf_export_uses_one_scale_for_every_block_on_a_page():
    content = make_pdf(["source"])
    exported = build_layout_pdf_export(
        content,
        [
            {
                "segment_id": "pdf:p1:b1",
                "page_number": 1,
                "text": "Long source",
                "translated_text": "Long translated sentence that must wrap across several lines",
                "source_kind": "text",
                "bbox": [72, 100, 172, 128],
                "font_size": 12.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {"leading": 14.0},
            },
            {
                "segment_id": "pdf:p1:b2",
                "page_number": 1,
                "text": "Short source",
                "translated_text": "Short",
                "source_kind": "text",
                "bbox": [72, 160, 172, 180],
                "font_size": 8.0,
                "color": "#000000",
                "alignment": "left",
                "metadata": {"leading": 9.0, "bold": True},
            },
        ],
        "en",
    )

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        spans = [
            span
            for block in document[0].get_text("dict")["blocks"]
            if block.get("type") == 0
            for line in block.get("lines", [])
            for span in line.get("spans", [])
        ]
        long_span = next(span for span in spans if "Long translated" in span["text"])
        short_span = next(span for span in spans if "Short" in span["text"])
        long_scale = long_span["size"] / 12.0
        short_scale = short_span["size"] / 8.0
        assert long_scale < 1.0
        assert abs(long_scale - short_scale) < 0.02
        assert "Bold" in short_span["font"]
    finally:
        document.close()


def test_unformatted_pdf_rejects_overflow_without_adding_pages():
    content = make_pdf(["source"])
    translated_text = "【第 1 页】\n" + "\n".join(
        f"第 {index} 行超长译文内容" for index in range(1, 181)
    )
    with pytest.raises(ValueError, match="LAYOUT_OVERFLOW"):
        build_pdf_export(content, translated_text, "zh")


def test_unformatted_pdf_uses_columns_for_dense_wide_pages():
    source = fitz.open()
    source.new_page(width=2384, height=1684)
    content = source.tobytes()
    source.close()
    translated_text = "【第 1 页】\n" + "\n".join(
        f"第 {index} 行密集工程说明" for index in range(1, 386)
    )

    exported = build_pdf_export(content, translated_text, "zh")

    document = fitz.open(stream=exported, filetype="pdf")
    try:
        assert document.page_count == 1
        extracted = document[0].get_text()
        assert "第 1 行密集工程说明" in extracted
        assert "第 385 行密集工程说明" in extracted
    finally:
        document.close()


def test_document_export_runs_outside_the_async_event_loop(monkeypatch, tmp_path):
    def slow_export(**kwargs):
        time.sleep(0.1)
        return ExportedDocument(
            filename="sample-译文.pdf",
            path=tmp_path / "sample.pdf",
            media_type="application/pdf",
        )

    monkeypatch.setattr(main_module, "create_document_export", slow_export)

    async def save_and_observe():
        task = asyncio.create_task(
            main_module._save_document_result(
                kind="document",
                source_text="source",
                target_language="zh",
                result=TranslationResult(
                    source_language="en",
                    translated_text="译文",
                    provider="fake",
                    warnings=[],
                ),
                filename="sample.pdf",
                source_content=make_pdf(["source"]),
            )
        )
        await asyncio.sleep(0.02)
        event_loop_remained_responsive = not task.done()
        result = await task
        return event_loop_remained_responsive, result

    remained_responsive, result = asyncio.run(save_and_observe())

    assert remained_responsive is True
    assert result.export_filename == "sample-译文.pdf"


def test_failed_document_job_keeps_source_and_diagnostics(monkeypatch):
    content = make_pdf(["A normal text document"])

    def rejected_pdf_page_count(_content):
        raise ValueError("PDF 页数预检失败：测试解析器异常")

    monkeypatch.setattr(main_module, "pdf_page_count", rejected_pdf_page_count)
    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document/jobs",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("diagnostic-source.pdf", content, "application/pdf")},
        )

    assert response.status_code == 422
    with main_module.database.connect() as connection:
        row = connection.execute(
            """
            SELECT * FROM document_attempts
            WHERE filename = ?
            ORDER BY created_at DESC LIMIT 1
            """,
            ("diagnostic-source.pdf",),
        ).fetchone()
    assert row is not None
    assert row["status"] == "failed"
    assert row["stage"] == "preflight"
    assert row["error_type"] == "ValueError"
    assert "测试解析器异常" in row["error_message"]
    assert "rejected_pdf_page_count" in row["error_trace"]
    assert row["content_bytes"] == len(content)
    assert row["content_sha256"]
    archived_source = settings.upload_source_directory / row["source_path"]
    assert archived_source.read_bytes() == content


def test_async_pdf_parse_failure_records_original_exception_chain(monkeypatch):
    content = make_pdf(["A normal text document"])

    def failing_iter_pdf_pages(_content):
        raise RuntimeError("低层 PDF 内容流解析失败")

    monkeypatch.setattr(main_module, "iter_pdf_pages", failing_iter_pdf_pages)
    with TestClient(app) as client:
        created = client.post(
            "/api/translate/document/jobs",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("async-parse-failure.pdf", content, "application/pdf")},
        )
        assert created.status_code == 202
        job_id = created.json()["job_id"]
        deadline = time.monotonic() + 5.0
        while time.monotonic() < deadline:
            status_response = client.get(f"/api/translate/document/jobs/{job_id}")
            body = status_response.json()
            if body["status"] != "processing":
                break
            time.sleep(0.01)

    assert body["status"] == "failed"
    assert body["error"] == "低层 PDF 内容流解析失败"
    row = main_module.database.get_document_attempt(job_id)
    assert row is not None
    assert row["status"] == "failed"
    assert row["stage"] == "translation"
    assert row["error_type"] == "RuntimeError"
    assert "failing_iter_pdf_pages" in row["error_trace"]
    assert (settings.upload_source_directory / row["source_path"]).read_bytes() == content


def test_unformatted_export_survives_formatted_export_failure(monkeypatch):
    def failed_export(**kwargs):
        raise ValueError("layout failed")

    monkeypatch.setattr(main_module, "create_document_export", failed_export)
    content = make_native_thai_pdf(["สวัสดี"])

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("failed-layout.pdf", content, "application/pdf")},
        )
        item_id = response.json()["id"]
        formatted = client.get(f"/api/history/{item_id}/export")
        unformatted = client.get(f"/api/history/{item_id}/export/unformatted")

    assert response.status_code == 200
    assert response.json()["export_filename"] is None
    assert "译文已生成，但格式化文件导出失败，请稍后重试" in response.json()["warnings"]
    assert formatted.status_code == 404
    assert unformatted.status_code == 200
    assert unformatted.headers["content-type"] == "application/pdf"


def test_unformatted_pdf_export_is_generated_on_demand_and_keeps_pages():
    content = make_native_thai_pdf(["สวัสดี", "ขอบคุณ"])

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("sample.pdf", content, "application/pdf")},
        )
        item_id = response.json()["id"]
        unformatted_path = settings.database_path.parent / "exports" / f"{item_id}-unformatted.pdf"
        assert not unformatted_path.exists()
        exported = client.get(f"/api/history/{item_id}/export/unformatted")
        cached = client.get(f"/api/history/{item_id}/export/unformatted")
        inline_export = client.get(
            f"/api/history/{item_id}/export/unformatted?inline=true",
            headers={"Range": "bytes=0-99"},
        )

    assert response.status_code == 200
    assert exported.status_code == 200
    assert cached.content == exported.content
    assert unformatted_path.is_file()
    assert exported.headers["content-disposition"].startswith(
        'attachment; filename="translation-unformatted.pdf"; filename*=UTF-8\'\''
    )
    assert "%E6%9C%AA%E6%8E%92%E7%89%88%E8%AF%91%E6%96%87.pdf" in exported.headers[
        "content-disposition"
    ]
    assert inline_export.status_code == 206
    assert inline_export.headers["content-disposition"].startswith("inline;")

    source_pdf = fitz.open(stream=content, filetype="pdf")
    translated_pdf = fitz.open(stream=exported.content, filetype="pdf")
    try:
        assert translated_pdf.page_count == source_pdf.page_count == 2
        for source_page, translated_page in zip(source_pdf, translated_pdf):
            assert translated_page.rect.width == source_page.rect.width
            assert translated_page.rect.height == source_page.rect.height
        assert "你好" in translated_pdf[0].get_text()
        assert "谢谢" not in translated_pdf[0].get_text()
        assert "谢谢" in translated_pdf[1].get_text()
        assert "你好" not in translated_pdf[1].get_text()
        assert translated_pdf[0].get_text("blocks")[0][0] < 60
    finally:
        translated_pdf.close()
        source_pdf.close()


def test_docx_export_uses_docx_and_contains_translation():
    document = Document()
    document.sections[0].top_margin = 720000
    document.sections[0].header.paragraphs[0].text = "Header text"
    paragraph = document.add_paragraph()
    source_run = paragraph.add_run("Hello document")
    source_run.font.size = Pt(18)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    stream = BytesIO()
    document.save(stream)

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={
                "file": (
                    "brief.docx",
                    stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        exported = client.get(
            f"/api/history/{response.json()['id']}/export"
        )

    assert response.status_code == 200
    assert response.json()["export_filename"] == "brief-译文.docx"
    assert exported.status_code == 200
    translated_document = Document(BytesIO(exported.content))
    assert "演示译文" in "\n".join(
        paragraph.text for paragraph in translated_document.paragraphs
    )
    translated_paragraph = next(
        paragraph
        for paragraph in translated_document.paragraphs
        if "演示译文" in paragraph.text
    )
    assert translated_paragraph.runs[0].font.size.pt == pytest.approx(18.0)
    assert abs(translated_document.sections[0].top_margin - 720000) < 200
    assert "Header text" in translated_document.sections[0].header.paragraphs[0].text
    assert len(translated_document.tables) == 1
    assert "Name" in translated_document.tables[0].cell(0, 0).text
    assert "Value" in translated_document.tables[0].cell(0, 1).text


def test_unformatted_docx_export_uses_continuous_text_layout():
    document = Document()
    document.sections[0].top_margin = 720000
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.add_paragraph("Hello document")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table value"
    stream = BytesIO()
    document.save(stream)

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={
                "file": (
                    "brief.docx",
                    stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        exported = client.get(
            f"/api/history/{response.json()['id']}/export/unformatted"
        )

    assert response.status_code == 200
    assert exported.status_code == 200
    assert exported.headers["content-disposition"].startswith(
        'attachment; filename="translation-unformatted.docx"; filename*=UTF-8\'\''
    )
    translated_document = Document(BytesIO(exported.content))
    assert "演示译文" in "\n".join(
        paragraph.text for paragraph in translated_document.paragraphs
    )
    assert abs(translated_document.sections[0].top_margin - 720000) < 200
    assert not translated_document.tables
    assert not translated_document.sections[0].header.paragraphs[0].text


def test_unformatted_export_rejects_pptx_and_text_documents():
    with TestClient(app) as client:
        pptx_response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={
                "file": (
                    "slides.pptx",
                    make_minimal_pptx(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        text_response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={"file": ("notes.txt", b"Hello document", "text/plain")},
        )
        pptx_export = client.get(
            f"/api/history/{pptx_response.json()['id']}/export/unformatted"
        )
        text_export = client.get(
            f"/api/history/{text_response.json()['id']}/export/unformatted"
        )

    assert pptx_response.status_code == 200
    assert text_response.status_code == 200
    assert pptx_export.status_code == 404
    assert text_export.status_code == 404


def test_pptx_export_replaces_text_in_place_and_enables_autofit():
    content = make_minimal_pptx()
    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={
                "file": (
                    "slides.pptx",
                    content,
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        exported = client.get(f"/api/history/{response.json()['id']}/export")

    assert response.status_code == 200
    assert response.json()["export_filename"] == "slides-译文.pptx"
    assert exported.status_code == 200
    with ZipFile(BytesIO(exported.content)) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "演示译文: Hello slide" in slide_xml
    assert "normAutofit" in slide_xml
    assert 'sz="2400"' in slide_xml


def test_knowledge_csv_import_uses_thai_key_and_latest_value_wins():
    csv_content = """source_language,target_language,source_text,translated_text
th,zh,มหาวิทยาลัยราชภัฏสวนสุนันทา,泰国川登喜皇家大学
th,zh,มหาวิทยาลัยราชภัฏสวนสุนันทา,泰国川登喜皇家大学
th,zh,มหาวิทยาลัยราชภัฏสวนสุนันทา,另一译法
zh,zh,无效,无效
""".encode("utf-8")

    with TestClient(app) as client:
        imported = client.post(
            "/api/knowledge/import",
            files={"file": ("terms.csv", csv_content, "text/csv")},
        )
        listed = client.get("/api/knowledge")
        translated = client.post(
            "/api/translate",
            json={
                "text": "มหาวิทยาลัยราชภัฏสวนสุนันทา",
                "source_language": "th",
                "target_language": "zh",
            },
        )

    assert imported.status_code == 200
    assert imported.json() == {
        "filename": "terms.csv",
        "total_rows": 4,
        "inserted": 1,
        "updated": 2,
        "invalid_rows": 1,
    }
    assert listed.status_code == 200
    assert listed.json()["total"] >= 1
    assert translated.status_code == 200
    assert translated.json()["translated_text"] == "另一译法"
    assert translated.json()["provider"] == "knowledge-base"


def test_knowledge_xlsx_import_uses_revision_columns_and_latest_thai_row():
    thai_text = "ระบบอาคารอัจฉริยะ"
    content = make_knowledge_xlsx(
        [
            [
                1,
                thai_text,
                "Smart building system",
                "Intelligent Building System",
                "智能楼宇系统（初译）",
                "智能建筑系统",
            ],
            [
                2,
                thai_text,
                "Final English Term",
                None,
                "最终中文初译",
                "最终中文术语",
            ],
            [3, None, "Invalid", None, "无效", None],
        ],
        header_row=3,
    )

    with TestClient(app) as client:
        imported = client.post(
            "/api/knowledge/import",
            files={
                "file": (
                    "revision-history.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        listed = client.get("/api/knowledge").json()

    assert imported.status_code == 200
    assert imported.json() == {
        "filename": "revision-history.xlsx",
        "total_rows": 3,
        "inserted": 1,
        "updated": 1,
        "invalid_rows": 1,
    }
    matching = next(item for item in listed["entries"] if item["thai_text"] == thai_text)
    assert matching["english_text"] == "Final English Term"
    assert matching["chinese_text"] == "最终中文术语"


def test_knowledge_xlsx_rejects_wrong_headers_and_unsupported_xls():
    content = make_knowledge_xlsx(
        [[1, "ระบบ", "System"]],
        headers=["NO", "Thai source", "Target"],
    )

    with TestClient(app) as client:
        malformed = client.post(
            "/api/knowledge/import",
            files={"file": ("wrong.xlsx", content)},
        )
        unsupported = client.post(
            "/api/knowledge/import",
            files={"file": ("legacy.xls", b"not-an-xls")},
        )

    assert malformed.status_code == 422
    assert "表头不符合模板" in malformed.json()["detail"]
    assert unsupported.status_code == 422
    assert unsupported.json()["detail"] == "知识库仅支持 CSV 或 XLSX 文件"


def test_knowledge_template_matches_revision_workbook_format():
    with TestClient(app) as client:
        response = client.get("/api/knowledge/template")
        imported = client.post(
            "/api/knowledge/import",
            files={"file": ("knowledge-template.xlsx", response.content)},
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["修订"]
        assert [cell.value for cell in workbook["修订"][1]] == KNOWLEDGE_XLSX_HEADERS
    finally:
        workbook.close()
    assert imported.status_code == 200
    assert imported.json() == {
        "filename": "knowledge-template.xlsx",
        "total_rows": 0,
        "inserted": 0,
        "updated": 0,
        "invalid_rows": 0,
    }


def test_manual_knowledge_requires_thai_and_overwrites_by_thai_key():
    thai_text = "ระบบท่อกำจัดปลวก"
    with TestClient(app) as client:
        missing_thai = client.post(
            "/api/knowledge/entries",
            json={"thai_text": "", "chinese_text": "管道式白蚁防治系统"},
        )
        missing_translation = client.post(
            "/api/knowledge/entries",
            json={"thai_text": thai_text},
        )
        created = client.post(
            "/api/knowledge/entries",
            json={
                "thai_text": thai_text,
                "chinese_text": "管道式白蚁防治系统",
                "english_text": "PIPE TREATMENT",
            },
        )
        updated = client.post(
            "/api/knowledge/entries",
            json={
                "thai_text": f"  {thai_text}  ",
                "chinese_text": "管道白蚁防治系统（新）",
            },
        )
        listed = client.get("/api/knowledge").json()

    assert missing_thai.status_code == 422
    assert missing_translation.status_code == 422
    assert created.status_code == 200
    assert updated.status_code == 200
    assert created.json()["id"] == updated.json()["id"]
    assert updated.json()["chinese_text"] == "管道白蚁防治系统（新）"
    assert updated.json()["english_text"] == "PIPE TREATMENT"
    matching = next(item for item in listed["entries"] if item["id"] == updated.json()["id"])
    assert matching["thai_text"] == thai_text
    assert main_module.database.find_exact_knowledge(
        "en", "zh", "PIPE TREATMENT"
    )["translated_text"] == "管道白蚁防治系统（新）"


def test_existing_directional_knowledge_is_migrated_to_thai_keyed_items(tmp_path):
    database_path = tmp_path / "legacy.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE knowledge_entries (
                id TEXT PRIMARY KEY,
                source_language TEXT NOT NULL,
                target_language TEXT NOT NULL,
                source_text TEXT NOT NULL,
                translated_text TEXT NOT NULL,
                source_normalized TEXT NOT NULL,
                translation_normalized TEXT NOT NULL,
                source_filename TEXT NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(source_language, target_language, source_normalized)
            )
            """
        )
        connection.execute(
            """
            INSERT INTO knowledge_entries VALUES
            ('legacy', 'th', 'zh', 'อาคาร', '建筑', 'อาคาร', '建筑',
             'legacy.csv', '2026-01-01T00:00:00+00:00')
            """
        )

    migrated = Database(database_path)
    migrated.initialize()

    listed = migrated.list_knowledge()
    assert listed["total"] == 1
    assert listed["entries"][0]["thai_text"] == "อาคาร"
    assert listed["entries"][0]["chinese_text"] == "建筑"
    assert migrated.find_exact_knowledge("zh", "th", "建筑")["translated_text"] == "อาคาร"


def test_knowledge_term_inside_longer_text_is_added_to_model_context(monkeypatch):
    main_module.database.upsert_knowledge_item(
        thai_text="ระบบท่อกำจัดปลวก",
        chinese_text="管道式白蚁防治系统",
        english_text="PIPE TREATMENT",
        source_filename="embedded.csv",
    )
    captured_context = []

    async def fake_translate(
        text,
        source_language,
        target_language,
        context,
        preserve_page_markers=False,
    ):
        captured_context.append(context)
        return TranslationResult(
            source_language="en",
            translated_text="采用管道式白蚁防治系统。",
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module.translator, "translate", fake_translate)

    with TestClient(app) as client:
        response = client.post(
            "/api/translate",
            json={
                "text": "The contractor shall use PIPE TREATMENT.",
                "source_language": "en",
                "target_language": "zh",
            },
        )

    assert response.status_code == 200
    assert response.json()["provider"] == "fake+knowledge-base"
    assert "PIPE TREATMENT => 管道式白蚁防治系统" in captured_context[0]
    assert any("知识库匹配项 1 条" in warning for warning in response.json()["warnings"])


def test_scanned_pdf_enters_visual_translation_and_can_still_be_rendered():
    content = make_pdf(["Invoice 2026-08-31"], image_only=True)

    parsed = parse_document("scan.pdf", content)
    rendered = render_pdf_pages(content, [1], 1200)

    assert parsed.ocr_required is True
    assert parsed.page_types == ["image"]
    assert parsed.page_profiles[0]["source_language"] == "auto"
    assert parsed.visual_pages == [1]
    assert parsed.page_count == 1
    assert parsed.text == ""
    assert len(rendered) == 1
    assert rendered[0].page_number == 1
    assert rendered[0].content.startswith(b"\x89PNG\r\n\x1a\n")


def test_vector_directory_tiles_use_three_full_height_columns():
    content = make_pdf(["directory"])
    rendered = render_pdf_tiles(content, [1], 1200, "columns")

    assert len(rendered) == 3
    assert all(item.clip is not None for item in rendered)
    assert all(item.clip[1] == 0.0 and item.clip[3] == 842.0 for item in rendered)
    assert rendered[0].clip[0] < rendered[0].clip[2] <= rendered[1].clip[0]


def test_scanned_pdf_translation_uses_visual_pipeline_without_native_text():
    content = make_pdf(["Page one", "Page two"], image_only=True)

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("scan.pdf", content, "application/pdf")},
        )

    assert response.status_code == 200
    assert response.json()["provider"].endswith(":layout")


def test_scanned_pdf_route_never_calls_cad_candidate_detection(monkeypatch):
    content = make_pdf(["scan"], image_only=True)
    rendered = RenderedPdfPage(page_number=1, content=b"png")

    monkeypatch.setattr(main_module, "render_pdf_pages", lambda *_args: [rendered])
    monkeypatch.setattr(
        main_module,
        "render_pdf_tiles",
        lambda *_args: pytest.fail("slide-sized scan was unnecessarily tiled"),
    )
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module.database, "list_knowledge_for_direction", lambda *_args, **_kwargs: []
    )
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        main_module,
        "ocr_image_text_blocks",
        lambda *_args: pytest.fail("scan route called CAD candidate detection"),
    )

    async def fake_translate_image_layout(*_args, **_kwargs):
        return "รายการ", TranslationResult(
            source_language="th",
            translated_text="项目",
            provider="vision",
            warnings=[],
            layout_segments=[
                {
                    "source_text": "BLACK LOCUST",
                    "translated_text": "BLACK LOCUST",
                    "bbox": [20, 20, 800, 160],
                },
                {
                    "source_text": "รายการ",
                    "translated_text": "项目",
                    "bbox": [100, 100, 400, 160],
                }
            ],
        )

    monkeypatch.setattr(
        main_module.translator,
        "translate_image_layout",
        fake_translate_image_layout,
    )

    page_results, _ = asyncio.run(
        main_module._translate_pdf_layout_pages(
            content,
            [1],
            "th",
            "zh",
            "",
            page_types={1: "image"},
        )
    )

    assert len(page_results) == 1
    assert len(page_results[0][3]) == 1
    assert page_results[0][3][0]["translated_text"] == "项目"


def test_scanned_translation_detects_preserved_handwriting_collision():
    source_rect = fitz.Rect(100, 100, 180, 116)
    handwriting = [fitz.Rect(150, 96, 190, 120)]

    assert _rect_overlap_area(source_rect, handwriting) > 0


def test_scanned_translation_has_no_false_handwriting_collision():
    source_rect = fitz.Rect(100, 100, 180, 116)

    assert _rect_overlap_area(source_rect, []) == 0


def test_scanned_translation_ignores_minor_handwriting_contact():
    source_rect = fitz.Rect(100, 100, 200, 120)
    handwriting = [fitz.Rect(195, 110, 210, 130)]

    assert _scan_ink_scale(source_rect, handwriting, 8.0, 1.0) == 1.0


def test_scanned_translation_keeps_colliding_text_readable():
    source_rect = fitz.Rect(100, 100, 200, 120)
    handwriting = [fitz.Rect(100, 100, 200, 120)]

    scale = _scan_ink_scale(source_rect, handwriting, 8.0, 0.8)

    assert 8.0 * 0.8 * scale >= 6.0
    assert scale < 1.0


def test_ocr_coverage_uses_page_coordinates_for_rendered_tiles():
    clip = (100.0, 100.0, 300.0, 300.0)
    visual = [{"bbox": [100.0, 100.0, 200.0, 120.0]}]
    covered = {
        "source_text": "รายการก่อสร้าง",
        "bbox": [0.0, 0.0, 500.0, 100.0],
    }
    uncovered = {
        "source_text": "วัสดุก่อสร้าง",
        "bbox": [500.0, 500.0, 900.0, 620.0],
    }

    selected = main_module._select_uncovered_ocr_blocks(
        [covered, uncovered], visual, clip
    )

    assert selected == [uncovered]


def test_text_layer_pdf_does_not_use_visual_ocr():
    content = make_pdf(["This is a selectable text PDF with enough content for translation."])

    parsed = parse_document("text.pdf", content, "en")
    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={"file": ("text.pdf", content, "application/pdf")},
        )

    assert parsed.ocr_required is False
    assert parsed.page_types == ["native_text"]
    assert parsed.visual_pages == []
    assert response.status_code == 200
    assert response.json()["source_language"] == "en"


def test_large_document_text_is_allowed_up_to_configured_limit(monkeypatch):
    document = ParsedDocument(
        filename="large.pdf",
        text="ก" * 394_278,
        page_count=206,
        warnings=[],
        page_texts=["ก"] * 206,
    )
    main_module._validate_document(document)

    monkeypatch.setattr(
        main_module,
        "settings",
        replace(settings, max_document_characters=394_277),
    )
    with pytest.raises(ValueError, match="394,277"):
        main_module._validate_document(document)


def test_layout_segments_are_batched_by_five_pages_and_safety_limits():
    segments = [
        DocumentSegment(
            segment_id=f"p{page}:b{index}",
            page_number=page,
            text="text",
        )
        for page in range(1, 8)
        for index in range(100)
    ]

    batches = main_module._build_document_segment_batches(segments)

    assert len(batches) == 4
    assert all(len(batch) <= main_module.LAYOUT_SEGMENT_MAX_ITEMS for batch in batches)
    assert all(
        len({segment.page_number for segment in batch})
        <= main_module.LAYOUT_SEGMENT_MAX_PAGES
        for batch in batches
    )
    assert all(
        sum(len(segment.text) for segment in batch)
        <= main_module.LAYOUT_SEGMENT_MAX_CHARACTERS
        for batch in batches
    )


def test_layout_segment_translation_reports_each_completed_batch(monkeypatch):
    calls = []
    progress = []
    segments = [
        DocumentSegment(
            segment_id=f"p{page}:b{index}",
            page_number=page,
            text=f"Page {page} block {index}",
        )
        for page in range(1, 8)
        for index in range(2)
    ]

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        calls.append(values)
        return SegmentTranslationResult(
            source_language="en",
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    translations, providers, warnings = asyncio.run(
        main_module._translate_document_segments(
            segments,
            "en",
            "zh",
            "",
            lambda completed, message: progress.append((completed, message)),
        )
    )

    assert len(calls) == 2
    assert len(translations) == len(segments)
    assert providers == ["fake", "fake"]
    assert sum(completed for completed, _ in progress) == 7
    assert len(progress) == 2
    assert all(message == "正在翻译并保留原排版" for _, message in progress)
    assert warnings[0] == "文案按最多 5 页分为 2 个请求"


@pytest.mark.parametrize("segment_count", [60, 208])
def test_dense_cad_text_batches_keep_full_page_context(monkeypatch, segment_count):
    calls = []
    segments = [
        DocumentSegment(
            segment_id=f"cad:p1:source:{index:03d}",
            page_number=1,
            text=f"อาคาร {index}",
            source_kind="outline-text",
        )
        for index in range(1, segment_count + 1)
    ]

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        calls.append((values, context))
        return SegmentTranslationResult(
            source_language=source_language,
            translations={key: f"建筑 {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    translations, _providers, _warnings = asyncio.run(
        main_module._translate_document_segments(segments, "th", "zh", "")
    )

    expected_batch_items = (
        main_module.CAD_BALANCED_TEXT_BATCH_ITEMS
        if segment_count <= main_module.CAD_BALANCED_TEXT_BATCH_MAX_PAGE_ITEMS
        else main_module.CAD_PARALLEL_TEXT_BATCH_ITEMS
    )
    assert len(calls) == (
        len(segments) + expected_batch_items - 1
    ) // expected_batch_items
    assert all(len(values) <= expected_batch_items for values, _ in calls)
    assert all(
        "[cad:p1:source:001] อาคาร 1" in context
        and f"[cad:p1:source:{segment_count:03d}] อาคาร {segment_count}" in context
        for _, context in calls
    )
    assert len(translations) == len(segments)


def test_malformed_layout_batch_is_split_instead_of_retrying_every_block(monkeypatch):
    calls = []
    segments = [
        DocumentSegment(
            segment_id=f"p1:b{index}",
            page_number=1,
            text=f"Block {index}",
        )
        for index in range(8)
    ]

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        calls.append(len(values))
        if len(values) > 2:
            raise RuntimeError("模型未完整返回全部文字块，请重试")
        return SegmentTranslationResult(
            source_language="en",
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    translations, _, warnings = asyncio.run(
        main_module._translate_document_segments(segments, "en", "zh", "")
    )

    assert len(translations) == 8
    assert calls == [8, 4, 4, 2, 2, 2, 2]
    assert any("已自动拆分并补发" in warning for warning in warnings)


def test_gateway_timeout_layout_batch_is_split_and_completed(monkeypatch):
    calls = []
    segments = [
        DocumentSegment(
            segment_id=f"p1:b{index}",
            page_number=1,
            text=f"Block {index}",
        )
        for index in range(8)
    ]

    async def fake_translate_segments(
        values, source_language, target_language, context, **_kwargs
    ):
        calls.append(len(values))
        if len(values) > 2:
            raise RuntimeError(
                '模型服务请求失败 (524): {"error":{"message":"openai_error"}}'
            )
        return SegmentTranslationResult(
            source_language="en",
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database,
        "find_matching_knowledge",
        lambda *args, **kwargs: [],
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    translations, _, warnings = asyncio.run(
        main_module._translate_document_segments(segments, "en", "zh", "")
    )

    assert len(translations) == 8
    assert calls == [8, 4, 4, 2, 2, 2, 2]
    assert any("已自动拆分并补发" in warning for warning in warnings)


def test_partial_layout_batch_retries_only_missing_ids_with_page_context(monkeypatch):
    calls = []
    segments = [
        DocumentSegment(
            segment_id=f"p1:b{index}",
            page_number=1,
            text=f"Block {index}",
        )
        for index in range(4)
    ]

    async def fake_translate_segments(values, _source, _target, context, **kwargs):
        calls.append((list(values), context, kwargs))
        if len(calls) == 1:
            return SegmentTranslationResult(
                source_language="en",
                translations={
                    "p1:b0": "translated: Block 0",
                    "p1:b1": "translated: Block 1",
                },
                provider="fake",
                warnings=[],
            )
        return SegmentTranslationResult(
            source_language="en",
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    translations, _, warnings = asyncio.run(
        main_module._translate_document_segments(segments, "en", "zh", "")
    )

    assert len(translations) == 4
    assert calls[0][0] == ["p1:b0", "p1:b1", "p1:b2", "p1:b3"]
    assert calls[0][2] == {"require_complete": False}
    assert calls[1][0] == ["p1:b2", "p1:b3"]
    assert "同一页的完整原文" in calls[1][1]
    assert "[p1:b0] Block 0" in calls[1][1]
    assert any("已自动拆分并补发" in warning for warning in warnings)


def test_mixed_pdf_routes_image_page_through_visual_pipeline():
    content = make_mixed_pdf()

    parsed = parse_document("mixed.pdf", content)
    assert parsed.ocr_required is True
    assert parsed.ocr_pages == [2]
    assert parsed.page_types == ["native_text", "image"]
    assert parsed.visual_pages == [2]
    assert parsed.page_texts[0].startswith("This is a selectable text page")
    assert parsed.page_texts[1] == ""

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("mixed.pdf", content, "application/pdf")},
        )

    assert response.status_code == 200
    assert response.json()["provider"].endswith(":layout")


def test_text_pages_are_sent_in_five_page_chunks_with_progress(monkeypatch):
    calls = []
    progress = []

    async def fake_translate(text, source_language, target_language, context, preserve_page_markers=False):
        calls.append(text)
        return TranslationResult(
            source_language="en",
            translated_text=text,
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module.translator, "translate", fake_translate)
    text_pages = [(page_number, f"Page {page_number} text") for page_number in range(1, 13)]

    outputs, warnings, providers, chunk_count = asyncio.run(
        main_module._translate_text_page_chunks(
            text_pages,
            "en",
            "zh",
            "",
            lambda completed, message: progress.append((completed, message)),
        )
    )

    assert chunk_count == 3
    assert len(calls) == 3
    assert all(text.count("【第") <= 5 for text in calls)
    assert [page_number for page_number, _, _ in outputs] == list(range(1, 13))
    assert sum(completed for completed, _ in progress) == 12
    assert all(message == "正在翻译文案页" for _, message in progress)
    assert warnings == []
    assert providers == ["fake", "fake", "fake"]


def test_text_pages_retry_individually_when_model_drops_page_markers(monkeypatch):
    calls = []
    progress = []

    async def fake_translate(
        text,
        source_language,
        target_language,
        context,
        preserve_page_markers=False,
    ):
        calls.append((text, preserve_page_markers))
        translated = "markers were dropped" if preserve_page_markers else f"translated: {text}"
        return TranslationResult(
            source_language="en",
            translated_text=translated,
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module.translator, "translate", fake_translate)
    text_pages = [(1, "page one"), (2, "page two"), (3, "page three")]

    outputs, warnings, providers, chunk_count = asyncio.run(
        main_module._translate_text_page_chunks(
            text_pages,
            "en",
            "zh",
            "",
            lambda completed, message: progress.append((completed, message)),
        )
    )

    assert chunk_count == 1
    assert len(calls) == 4
    assert calls[0][1] is True
    assert all(preserve_page_markers is False for _, preserve_page_markers in calls[1:])
    assert [page_number for page_number, _, _ in outputs] == [1, 2, 3]
    assert [result.translated_text for _, _, result in outputs] == [
        "translated: page one",
        "translated: page two",
        "translated: page three",
    ]
    assert sum(completed for completed, _ in progress) == 3
    assert all(message == "正在按单页校正页码" for _, message in progress)
    assert warnings == ["第 1-3 页翻译未保留页码标记，已自动按单页重试以保持页码对应"]
    assert providers == ["fake", "fake", "fake"]


def test_document_job_reports_progress_and_returns_result():
    content = make_native_thai_pdf(["สวัสดี", "ขอบคุณ"])

    with TestClient(app) as client:
        created = client.post(
            "/api/translate/document/jobs",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("mixed.pdf", content, "application/pdf")},
        )

        assert created.status_code == 202
        job = created.json()
        assert job["total_pages"] == 2
        assert 0 <= job["completed_pages"] <= 2

        deadline = time.monotonic() + 5.0
        while time.monotonic() < deadline:
            status_response = client.get(
                f"/api/translate/document/jobs/{job['job_id']}"
            )
            assert status_response.status_code == 200
            job = status_response.json()
            if job["status"] != "processing":
                break
            time.sleep(0.01)

    assert job["status"] == "completed"
    assert job["stage"] == "completed"
    assert job["completed_pages"] == 2
    assert job["progress"] == 100
    assert job["message"] == "翻译完成"
    assert job["result"]["kind"] == "document"
    assert job["result"]["provider"] == "demo:layout"


def test_pdf_pipeline_starts_translation_before_remaining_pages_are_parsed(
    monkeypatch,
):
    translation_started = threading.Event()
    parser_continued = threading.Event()

    def fake_iter_pdf_pages(_):
        for page_number in range(1, 6):
            text = f"Page {page_number}"
            yield ParsedPdfPage(
                page_number=page_number,
                text=text,
                segments=[
                    DocumentSegment(
                        segment_id=f"pdf:p{page_number}:b1",
                        page_number=page_number,
                        text=text,
                    )
                ],
            )
        assert translation_started.wait(1)
        parser_continued.set()
        yield ParsedPdfPage(
            page_number=6,
            text="Page 6",
            segments=[
                DocumentSegment(
                    segment_id="pdf:p6:b1",
                    page_number=6,
                    text="Page 6",
                )
            ],
        )

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        translation_started.set()
        return SegmentTranslationResult(
            source_language=source_language,
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "iter_pdf_pages", fake_iter_pdf_pages)
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            b"pdf",
            "stream.pdf",
            6,
            "en",
            "zh",
            "",
        )
    )

    assert parser_continued.is_set()
    assert source_text.count("【第") == 6
    assert result.translated_text.count("translated: Page") == 6
    assert result.provider == "fake:layout"


def test_pdf_pipeline_uses_prepared_source_from_streaming_parser(monkeypatch):
    prepared = b"prepared-pdf"

    def fake_iter_pdf_pages(_content, _source_language, prepared_callback):
        yield ParsedPdfPage(
            page_number=1,
            text="English",
            segments=[
                DocumentSegment(
                    segment_id="pdf:p1:s1:o1:t0",
                    page_number=1,
                    text="English",
                    metadata={
                        "engine": "content-stream",
                        "native_pdf_version": 4,
                        "code_refs": [
                            {
                                "stream_xref": 1,
                                "operation_index": 1,
                                "array_index": -1,
                                "token_index": 0,
                            }
                        ],
                    },
                )
            ],
        )
        prepared_callback(prepared)

    async def fake_translate_segments(
        values, source_language, target_language, context, **_kwargs
    ):
        return SegmentTranslationResult(
            source_language=source_language,
            translations={key: "英译" for key in values},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "iter_pdf_pages", fake_iter_pdf_pages)
    monkeypatch.setattr(
        main_module,
        "prepare_native_pdf_source",
        lambda *_args: pytest.fail("不应重复解析已准备的 PDF"),
    )
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    _, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            b"pdf",
            "stream.pdf",
            1,
            "en",
            "zh",
            "",
        )
    )

    assert result.prepared_pdf_content == prepared


def test_pdf_pipeline_uses_the_same_five_page_batches(monkeypatch):
    pages = []
    segments = []
    for page_number in range(1, 13):
        page_segments = [
            DocumentSegment(
                segment_id=f"pdf:p{page_number}:b{block_number}",
                page_number=page_number,
                text=f"Page {page_number} block {block_number}",
            )
            for block_number in range(1, 3)
        ]
        pages.append(
            ParsedPdfPage(
                page_number=page_number,
                text="\n".join(segment.text for segment in page_segments),
                segments=page_segments,
            )
        )
        segments.extend(page_segments)
    captured_batches = []

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        captured_batches.append(values)
        return SegmentTranslationResult(
            source_language=source_language,
            translations={key: f"translated: {value}" for key, value in values.items()},
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter(pages))
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    asyncio.run(
        main_module._translate_pdf_document_pipeline(
            b"pdf",
            "stream.pdf",
            len(pages),
            "en",
            "zh",
            "",
        )
    )

    expected_batches = [
        {segment.segment_id: segment.text for segment in batch}
        for batch in main_module._build_document_segment_batches(segments)
    ]
    assert captured_batches == expected_batches


def test_pdf_pipeline_uses_detected_source_language(monkeypatch):
    pages = [
        ParsedPdfPage(
            page_number=page_number,
            text="English",
            segments=[
                DocumentSegment(
                    segment_id=f"pdf:p{page_number}:b1",
                    page_number=page_number,
                    text="English",
                )
            ],
        )
        for page_number in range(1, 6)
    ]
    requested_languages = []

    async def fake_translate_segments(values, source_language, target_language, context, **_kwargs):
        requested_languages.append(source_language)
        return SegmentTranslationResult(
            source_language=source_language,
            translations={
                key: f"{source_language}: {value}" for key, value in values.items()
            },
            provider="fake",
            warnings=[],
        )

    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter(pages))
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *args: {}
    )
    monkeypatch.setattr(
        main_module.database, "find_matching_knowledge", lambda *args, **kwargs: []
    )
    monkeypatch.setattr(
        main_module.translator, "translate_segments", fake_translate_segments
    )

    _, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            b"pdf",
            "stream.pdf",
            len(pages),
            "auto",
            "zh",
            "",
        )
    )

    assert requested_languages
    assert set(requested_languages) == {"en"}
    assert result.translated_text.count("en: English") == 5
    assert result.source_language == "en"


def test_unknown_document_job_returns_404():
    with TestClient(app) as client:
        response = client.get("/api/translate/document/jobs/not-found")

    assert response.status_code == 404


def test_numeric_normalization_avoids_thai_digit_and_page_marker_false_positives():
    translated = normalize_translation_text("总面积为๓,๓๖๔平方米&#x20;至佛历๒๕๖๗年")

    assert translated == "总面积为3,364平方米 至佛历2567年"
    assert quality_checks(
        "【第 ๒๑ 页】总面积为๓,๓๖๔平方米",
        "【第 21 页】总面积为3,364平方米",
    ) == []
    assert quality_checks("付款比例为3.30%", "付款比例为3.3%") == []
    assert quality_checks("付款比例为3.30", "付款比例为3.30%") == []


def test_segment_translation_sends_complete_mixed_line_without_placeholders():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.return_value = (
        {"line-1": "ABC 建筑 123 / A-01"},
        "test",
    )
    source = {"line-1": "ABC อาคาร 123 / A-01"}

    result = asyncio.run(service.translate_segments(source, "th", "zh"))

    assert service.provider.translate_segments.await_args.args[0] == source
    assert result.translations == {"line-1": "ABC 建筑 123 / A-01"}


def test_segment_translation_treats_empty_output_as_missing():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.return_value = (
        {"line-1": "", "line-2": "已翻译"},
        "test",
    )
    source = {"line-1": "Missing", "line-2": "Translated"}

    partial = asyncio.run(
        service.translate_segments(source, "en", "zh", require_complete=False)
    )
    assert partial.translations == {"line-2": "已翻译"}

    with pytest.raises(RuntimeError, match="ID_MISMATCH.*译文为空"):
        asyncio.run(service.translate_segments(source, "en", "zh"))


def test_segment_translation_rejects_id_mismatch_and_romanizes_stubborn_residual():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.return_value = ({"line-1": "建筑"}, "test")

    result = asyncio.run(
        service.translate_segments({"line-1": "ABC อาคาร 123"}, "th", "zh")
    )
    assert result.translations == {"line-1": "建筑"}

    service.provider.translate_segments.return_value = ({"wrong-id": "建筑"}, "test")
    with pytest.raises(RuntimeError, match="ID_MISMATCH"):
        asyncio.run(
            service.translate_segments({"line-1": "ABC อาคาร 123"}, "th", "zh")
        )

    service.provider.translate_segments.return_value = ({"line-1": "ABC อาคาร"}, "test")
    service.provider.translate_segments.side_effect = None
    service.provider.translate_segments.return_value = ({"line-1": "ABC อาคาร"}, "test")
    service.provider.translate.return_value = ("ABC 建筑 123", "test")
    retried = asyncio.run(
        service.translate_segments({"line-1": "ABC อาคาร 123"}, "th", "zh")
    )
    assert retried.translations == {"line-1": "ABC 建筑 123"}

    service.provider.translate.return_value = ("ABC อาคาร", "test")
    repaired = asyncio.run(
        service.translate_segments({"line-1": "ABC อาคาร 123"}, "th", "zh")
    )
    assert not any("\u0E00" <= char <= "\u0E7F" for char in repaired.translations["line-1"])
    assert repaired.warnings == [
        "模型多次复核后仍保留泰文，已将 1 个文字块的剩余片段按泰语拉丁转写"
    ]


def test_segment_translation_romanizes_isolated_thai_cad_letter():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.return_value = (
        {"ID006": "ฏ（建筑物"},
        "vision",
    )
    service.provider.translate.return_value = ("ฏ（建筑物", "text")

    result = asyncio.run(
        service.translate_segments({"ID006": "ฏ (อาคาร"}, "th", "zh")
    )

    assert result.translations == {"ID006": "T（建筑物"}


def test_segment_translation_retries_complete_unit_until_thai_is_removed():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.return_value = (
        {"line-1": "特殊混凝土板墙 15 โต๊ะเต่า"},
        "vision",
    )
    service.provider.translate.side_effect = [
        ("特殊混凝土板墙 15 โต๊ะเต่า", "text"),
        ("特殊混凝土板墙，厚度 15，按产品标准安装", "text"),
    ]

    result = asyncio.run(
        service.translate_segments(
            {"line-1": "ผนังแผงคอนกรีตแบบพิเศษหนา 15 โต๊ะเต่า"},
            "th",
            "zh",
        )
    )

    assert result.translations == {"line-1": "特殊混凝土板墙，厚度 15，按产品标准安装"}
    assert service.provider.translate.call_count == 2
    assert all(
        call.args[0] == "ผนังแผงคอนกรีตแบบพิเศษหนา 15 โต๊ะเต่า"
        for call in service.provider.translate.await_args_list
    )


def test_segment_translation_repairs_multiple_residuals_in_one_structured_call():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_segments.side_effect = [
        (
            {"line-1": "墙体 ผนัง", "line-2": "地面 พื้น"},
            "initial",
        ),
        (
            {"line-1": "墙体", "line-2": "地面"},
            "bulk-repair",
        ),
    ]

    result = asyncio.run(
        service.translate_segments(
            {"line-1": "ผนัง", "line-2": "พื้น"}, "th", "zh"
        )
    )

    assert result.translations == {"line-1": "墙体", "line-2": "地面"}
    assert service.provider.translate_segments.await_count == 2
    assert service.provider.translate.await_count == 0


def test_indexed_image_translation_keeps_complete_mixed_line_and_validates_ids():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "ID001",
                "source_text": "W6 ผนัง 123",
                "translated_text": "W6 墙体 123",
            },
            {
                "id": "ID002",
                "source_text": "PURE ENGLISH",
                "translated_text": "PURE ENGLISH",
            },
        ],
        "vision",
    )

    items, route = asyncio.run(
        service.translate_indexed_image_lines(
            b"image", "image/png", ["ID001", "ID002"], "zh"
        )
    )

    assert route == "vision"
    assert items == [
        {
            "id": "ID001",
            "source_text": "W6 ผนัง 123",
            "translated_text": "W6 墙体 123",
        },
    ]

    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "ID001",
                "source_text": "W6 ผนัง 123",
                "translated_text": "W6 墙体 123",
            },
            {
                "id": "ID001",
                "source_text": "W6 ผนัง 123",
                "translated_text": "重复项",
            },
        ],
        "vision",
    )
    deduplicated, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image", "image/png", ["ID001"], "zh"
        )
    )
    assert deduplicated[0]["translated_text"] == "W6 墙体 123"

    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "UNKNOWN",
                "source_text": "ผนัง",
                "translated_text": "墙体",
            }
        ],
        "vision",
    )
    with pytest.raises(RuntimeError, match="ID_MISMATCH"):
        asyncio.run(
            service.translate_indexed_image_lines(
                b"image", "image/png", ["ID001"], "zh"
            )
        )


def test_indexed_image_translation_retries_and_rejects_missing_ids():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.side_effect = [
        (
            [
                {
                    "id": "ID001",
                    "source_text": "W6 ผนัง 123",
                    "translated_text": "W6 墙体 123",
                }
            ],
            "vision",
        ),
        (
            [
                {
                    "id": "ID002",
                    "source_text": "F7 พื้น",
                    "translated_text": "F7 地面",
                }
            ],
            "vision",
        ),
    ]

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image", "image/png", ["ID001", "ID002"], "zh"
        )
    )
    assert [item["id"] for item in items] == ["ID001", "ID002"]
    assert service.provider.translate_indexed_image_lines.call_count == 2

    service.provider.translate_indexed_image_lines.reset_mock()
    service.provider.translate_indexed_image_lines.side_effect = None
    service.provider.translate_indexed_image_lines.return_value = ([], "vision")
    with pytest.raises(RuntimeError, match="ID_MISMATCH: CAD 索引图缺少 2 个 ID"):
        asyncio.run(
            service.translate_indexed_image_lines(
                b"image", "image/png", ["ID001", "ID002"], "zh"
            )
        )

    partial, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001", "ID002"],
            "zh",
            require_complete=False,
            include_non_thai=True,
        )
    )
    assert partial == []


def test_indexed_image_translation_retries_malformed_model_json():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.side_effect = [
        RuntimeError("CAD 索引图翻译结果格式不正确，请重试"),
        (
            [
                {
                    "id": "ID001",
                    "source_text": "W6 ผนัง 123",
                    "translated_text": "W6 墙体 123",
                }
            ],
            "vision",
        ),
    ]

    items, route = asyncio.run(
        service.translate_indexed_image_lines(
            b"image", "image/png", ["ID001"], "zh"
        )
    )

    assert service.provider.translate_indexed_image_lines.call_count == 2
    assert route == "vision"
    assert items[0]["translated_text"] == "W6 墙体 123"


def test_indexed_image_translation_passes_noisy_source_hints_to_vision():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "ID001",
                "source_text": "W6 ผนัง 123",
                "translated_text": "W6 墙体 123",
            }
        ],
        "vision",
    )

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001"],
            "zh",
            expected_sources={"ID001": "W6 ผ นั ง 123"},
        )
    )

    assert items[0]["translated_text"] == "W6 墙体 123"
    assert service.provider.translate_indexed_image_lines.await_args.kwargs[
        "expected_sources"
    ] == {"ID001": "W6 ผ นั ง 123"}


def test_openai_indexed_image_prompt_marks_ocr_hints_as_noisy():
    provider = OpenAIProvider(settings)
    provider._generate_with_image = AsyncMock(
        return_value=(
            '{"items":[{"id":"ID001","source_text":"ผนัง",'
            '"translated_text":"墙体"}]}',
            "openai:responses",
        )
    )

    items, _ = asyncio.run(
        provider.translate_indexed_image_lines(
            b"image",
            "image/png",
            "zh",
            "",
            expected_sources={"ID001": "ผ นั ง"},
        )
    )

    instructions = provider._generate_with_image.await_args.args[0]
    assert "有噪声识别提示" in instructions
    assert '"ID001": "ผ นั ง"' in instructions
    assert items[0]["translated_text"] == "墙体"


def test_openai_grouped_indexed_reader_keeps_each_image_at_high_detail():
    provider = OpenAIProvider(settings)
    provider._generate_with_images = AsyncMock(
        return_value=(
            '{"items":[{"id":"ID0001","source_text":"ผนัง"},'
            '{"id":"ID0002","source_text":"ประตู"}]}',
            "openai:responses",
        )
    )

    items, _ = asyncio.run(
        provider.read_indexed_image_line_group(
            [b"sheet-1", b"sheet-2"],
            "image/png",
            "",
            expected_sources={"ID0001": "ผ นั ง"},
        )
    )

    call = provider._generate_with_images.await_args
    assert len(call.args[1]) == 2
    assert call.kwargs["detail"] == "high"
    assert "全局唯一" in call.args[0]
    assert [item["id"] for item in items] == ["ID0001", "ID0002"]


def test_openai_provider_reuses_and_closes_http_client_per_loop(monkeypatch):
    clients = []

    class FakeAsyncClient:
        def __init__(self, **kwargs):
            self.kwargs = kwargs
            self.is_closed = False
            self.posts = []
            clients.append(self)

        async def post(self, url, **kwargs):
            self.posts.append((url, kwargs))
            return httpx.Response(200, json={"output_text": "ok"})

        async def aclose(self):
            self.is_closed = True

    monkeypatch.setattr(
        "app.services.translator.httpx.AsyncClient",
        FakeAsyncClient,
    )
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_base_url="https://gateway.example/v1",
        )
    )

    async def exercise_provider():
        await provider._post("/responses", {"input": "one"})
        await provider._post("/responses", {"input": "two"})
        assert len(clients) == 1
        assert len(clients[0].posts) == 2
        assert clients[0].kwargs["trust_env"] is False
        assert clients[0].kwargs["limits"].max_connections >= 4
        await provider.aclose()

    asyncio.run(exercise_provider())

    assert clients[0].is_closed is True
    assert provider._clients_by_loop == {}
    assert provider._semaphores_by_loop == {}


def test_cad_source_hint_filter_keeps_words_and_rejects_line_noise():
    assert main_module._cad_source_hint_is_reliable(
        {
            "source_hint": "W6 ผนัง 123",
            "source_confidence": 82.0,
        }
    )
    assert not main_module._cad_source_hint_is_reliable(
        {
            "source_hint": "งจจงจ",
            "source_confidence": 64.0,
        }
    )
    assert not main_module._cad_source_hint_is_reliable(
        {
            "source_hint": "ผนัง",
            "source_confidence": 48.0,
        }
    )


def test_repeated_cad_source_hint_recovers_only_close_complete_lines():
    canonical = [
        "ST-05 ดูแบบขยาย A6-05",
        "ห้องไฟฟ้า",
    ]

    assert main_module._match_repeated_cad_source_hint(
        "ST05 ดแบบขยาย A605",
        canonical,
    ) == "ST-05 ดูแบบขยาย A6-05"
    assert main_module._match_repeated_cad_source_hint(
        "งจจงจ",
        canonical,
    ) == ""
    assert main_module._match_repeated_cad_source_hint(
        "ขึ้น",
        canonical,
    ) == ""


def test_indexed_image_translation_accepts_thai_after_two_consistent_readings():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.side_effect = [
        (
            [
                {
                    "id": "ID001",
                    "source_text": "PURE ENGLISH",
                    "translated_text": "PURE ENGLISH",
                }
            ],
            "vision",
        ),
        (
            [
                {
                    "id": "ID001",
                    "source_text": "W6 ผนัง 123",
                    "translated_text": "W6 墙体 123",
                }
            ],
            "vision",
        ),
        (
            [
                {
                    "id": "ID001",
                    "source_text": "W6 ผนัง 123",
                    "translated_text": "W6 墙体 123",
                }
            ],
            "vision",
        ),
    ]

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001"],
            "zh",
            verify_thai_source=True,
        )
    )

    assert service.provider.translate_indexed_image_lines.call_count == 3
    assert items[0]["translated_text"] == "W6 墙体 123"


def test_indexed_image_translation_rejects_inconsistent_thai_readings():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.side_effect = [
        ([{"id": "ID001", "source_text": "กกก", "translated_text": "甲"}], "vision"),
        ([{"id": "ID001", "source_text": "ขขข", "translated_text": "乙"}], "vision"),
        ([{"id": "ID001", "source_text": "คคค", "translated_text": "丙"}], "vision"),
    ]

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001"],
            "zh",
            require_complete=False,
            verify_thai_source=True,
        )
    )

    assert items == []


def test_indexed_image_translation_accepts_confirmed_non_thai_false_positive():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    false_positive = [
        {
            "id": "ID001",
            "source_text": "PMC",
            "translated_text": "PMC",
        }
    ]
    service.provider.translate_indexed_image_lines.return_value = (
        false_positive,
        "vision",
    )

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001"],
            "zh",
            include_non_thai=True,
            verify_thai_source=True,
        )
    )

    assert service.provider.translate_indexed_image_lines.call_count == 2
    assert items == false_positive


def test_pdf_pipeline_uses_indexed_gpt_images_for_dense_cad(monkeypatch):
    content = make_pdf(["CAD"])
    native_segment = DocumentSegment(
        segment_id="pdf:p1:n1",
        page_number=1,
        text="ชื่อโครงการ",
    )
    page = ParsedPdfPage(
        page_number=1,
        text=native_segment.text,
        segments=[native_segment],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "อาคาร",
        "source_confidence": 99.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module.database, "find_exact_knowledge_many", lambda *_args: {}
    )
    monkeypatch.setattr(
        main_module.database,
        "find_matching_knowledge",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"png", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )

    async def fake_read_indexed(*_args, **kwargs):
        assert kwargs["expected_sources"] == {"ID001": "อาคาร"}
        return (
            [
                {
                    "id": "ID001",
                    "source_text": "ABC อาคาร 123",
                }
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_read_indexed,
    )

    async def fake_translate_segments(segments, *_args, **_kwargs):
        assert [segment.text for segment in segments] == ["ABC อาคาร 123"]
        return {segments[0].segment_id: "ABC 建筑 123"}, ["text"], []

    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        fake_translate_segments,
    )

    async def fake_translate_native(batch, *_args, **_kwargs):
        return {batch[0].segment_id: "项目名称"}, ["text"], []

    monkeypatch.setattr(
        main_module,
        "_translate_document_segment_batch",
        fake_translate_native,
    )

    progress = []
    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content,
            "cad.pdf",
            1,
            "th",
            "zh",
            "",
            lambda completed, message: progress.append((completed, message)),
        )
    )

    assert "ABC อาคาร 123" in source_text
    assert result.provider == "text+vision:layout"
    visual_segment = next(
        segment
        for segment in result.layout_segments
        if segment.get("metadata", {}).get("ocr_provider") == "gpt-indexed-source"
    )
    assert visual_segment["translated_text"] == "ABC 建筑 123"
    assert all(completed == 0 for completed, _message in progress[:-1])
    assert any("定位 CAD 文字" in message for _completed, message in progress)
    assert any("读取 CAD 原文" in message for _completed, message in progress)
    assert any("复核 CAD 漏检文字" in message for _completed, message in progress)
    assert any("翻译并回写 CAD 文字" in message for _completed, message in progress)


def test_dense_cad_group_hedges_only_after_primary_is_slow(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidates = [
        {
            "bbox": (40.0, 60.0 + index * 30.0, 210.0, 82.0 + index * 30.0),
            "rotation": 0,
            "vertical": False,
            "source_hint": "อาคาร",
            "source_confidence": 99.0,
        }
        for index in range(2)
    ]

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(main_module, "CAD_INDEXED_IMAGES_PER_REQUEST", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_HEDGE_DELAY_SECONDS", 0.01)
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"one", "entries": {"ID0001": candidates[0]}},
            {"content": b"two", "entries": {"ID0002": candidates[1]}},
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    calls = 0

    async def fake_group_reader(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            await asyncio.sleep(0.05)
        return (
            [
                {"id": "ID0001", "source_text": "อาคาร 1"},
                {"id": "ID0002", "source_text": "อาคาร 2"},
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_line_group",
        fake_group_reader,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert calls == 2
    assert "อาคาร 1" in source_text
    assert len(result.layout_segments) == 2


def test_dense_cad_single_sheet_read_uses_same_slow_request_hedge(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "อาคาร",
        "source_confidence": 99.0,
    }
    events = []

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(main_module, "CAD_INDEXED_IMAGES_PER_REQUEST", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_MODEL_CONCURRENCY", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_HEDGE_DELAY_SECONDS", 0.01)
    monkeypatch.setattr(
        main_module,
        "_log_document_event",
        lambda event, **details: events.append((event, details)),
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"one", "entries": {"ID0001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    calls = 0

    async def fake_single_reader(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            await asyncio.sleep(0.05)
        return ([{"id": "ID0001", "source_text": "อาคาร"}], "vision")

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_single_reader,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    _, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert calls == 2
    assert len(result.layout_segments) == 1
    assert [
        details["image_count"]
        for event, details in events
        if event == "cad_indexed_vision_hedge_started"
    ] == [1]


def test_dense_cad_group_hedges_when_slot_frees_after_delay(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    events = []

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(main_module, "CAD_INDEXED_IMAGES_PER_REQUEST", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_MODEL_CONCURRENCY", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_HEDGE_DELAY_SECONDS", 0.01)
    monkeypatch.setattr(
        main_module,
        "_log_document_event",
        lambda event, **details: events.append((event, details)),
    )

    def candidate(index):
        return {
            "bbox": (40.0, 30.0 * index, 210.0, 30.0 * index + 22.0),
            "rotation": 0,
            "vertical": False,
            "source_hint": f"อาคาร {index}",
            "source_confidence": 99.0,
        }

    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"slow-a", "entries": {"ID0001": candidate(1)}},
            {"content": b"slow-b", "entries": {"ID0002": candidate(2)}},
            {"content": b"fast-a", "entries": {"ID0003": candidate(3)}},
            {"content": b"fast-b", "entries": {"ID0004": candidate(4)}},
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    calls = {"slow": 0, "fast": 0}

    async def fake_group_reader(images, _mime, expected_ids, *_args, **_kwargs):
        group = "slow" if images[0].startswith(b"slow") else "fast"
        calls[group] += 1
        if group == "slow" and calls[group] == 1:
            await asyncio.sleep(0.08)
        elif group == "fast":
            await asyncio.sleep(0.04)
        return (
            [
                {"id": item_id, "source_text": f"อาคาร {item_id}"}
                for item_id in expected_ids
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_line_group",
        fake_group_reader,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    _, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert calls == {"slow": 2, "fast": 1}
    assert len(result.layout_segments) == 4
    assert [
        details["group_number"]
        for event, details in events
        if event == "cad_indexed_vision_hedge_started"
    ] == [1]


def test_dense_cad_group_does_not_hedge_while_primary_waits_for_slot(monkeypatch):
    content = make_pdf(["CAD 1", "CAD 2"])
    pages = [
        ParsedPdfPage(
            page_number=page_number,
            text="",
            segments=[],
            page_type="vector",
            profile={"drawing_count": 50_000, "visual_required": True},
        )
        for page_number in (1, 2)
    ]
    events = []

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter(pages))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(main_module, "CAD_INDEXED_IMAGES_PER_REQUEST", 2)
    monkeypatch.setattr(main_module, "CAD_INDEXED_MODEL_CONCURRENCY", 1)
    monkeypatch.setattr(main_module, "CAD_INDEXED_HEDGE_DELAY_SECONDS", 0.02)
    monkeypatch.setattr(
        main_module,
        "_log_document_event",
        lambda event, **details: events.append((event, details)),
    )

    def fake_prepare(page, *_args, **_kwargs):
        page_label = "1" if "CAD 1" in page.get_text() else "2"
        return [
            {
                "content": f"page-{page_label}-a".encode(),
                "entries": {
                    "ID0001": {
                        "bbox": (10.0, 10.0, 80.0, 24.0),
                        "rotation": 0,
                        "vertical": False,
                        "source_hint": "อาคาร 1",
                        "source_confidence": 99.0,
                    }
                },
            },
            {
                "content": f"page-{page_label}-b".encode(),
                "entries": {
                    "ID0002": {
                        "bbox": (10.0, 30.0, 80.0, 44.0),
                        "rotation": 0,
                        "vertical": False,
                        "source_hint": "อาคาร 2",
                        "source_confidence": 99.0,
                    }
                },
            },
        ]

    monkeypatch.setattr(
        main_module, "prepare_dense_cad_translation_sheets", fake_prepare
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )

    async def fake_group_reader(images, *_args, **_kwargs):
        if images[0].startswith(b"page-1"):
            await asyncio.sleep(0.08)
        return (
            [
                {"id": "ID0001", "source_text": "อาคาร 1"},
                {"id": "ID0002", "source_text": "อาคาร 2"},
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_line_group",
        fake_group_reader,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 2, "th", "zh", ""
        )
    )

    hedge_pages = [
        details["page_number"]
        for event, details in events
        if event == "cad_indexed_vision_hedge_started"
    ]
    assert hedge_pages == []


def test_dense_cad_finishes_tesseract_sheet_build_before_paddle_detection(
    monkeypatch,
):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    tesseract_finished = threading.Event()
    page_stream_finished = threading.Event()

    def fake_iter_pages(*_args):
        yield page
        page_stream_finished.set()

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", fake_iter_pages)
    monkeypatch.setattr(main_module.translator, "provider", object())

    def fake_detect(*_args, **kwargs):
        assert kwargs["defer_existing_filter"] is True
        assert kwargs["selective_recognition"] is True
        assert kwargs["existing_bboxes"] == []
        assert tesseract_finished.is_set()
        return []

    def fake_prepare(*_args, **_kwargs):
        assert page_stream_finished.is_set()
        assert not tesseract_finished.is_set()
        tesseract_finished.set()
        return []

    monkeypatch.setattr(
        main_module, "detect_dense_cad_paddle_candidates", fake_detect
    )
    monkeypatch.setattr(
        main_module, "prepare_dense_cad_translation_sheets", fake_prepare
    )
    monkeypatch.setattr(
        main_module,
        "_build_layout_translation_result",
        lambda *_args: (
            "",
            TranslationResult(
                source_language="th",
                translated_text="",
                provider="test",
                warnings=[],
            ),
        ),
    )

    asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )
    assert tesseract_finished.is_set()


def test_dense_cad_reuses_one_paddle_worker_across_pages(monkeypatch):
    content = make_pdf(["CAD 1", "CAD 2"])
    pages = [
        ParsedPdfPage(
            page_number=page_number,
            text="",
            segments=[],
            page_type="vector",
            profile={"drawing_count": 50_000, "visual_required": True},
        )
        for page_number in (1, 2)
    ]
    worker_threads = []

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter(pages))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [],
    )

    def fake_detect(*_args, **_kwargs):
        worker_threads.append(threading.get_ident())
        return []

    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        fake_detect,
    )

    for _ in range(2):
        with pytest.raises(RuntimeError, match="没有可翻译的文字"):
            asyncio.run(
                main_module._translate_pdf_document_pipeline(
                    content, "cad.pdf", 2, "th", "zh", ""
                )
            )

    assert len(worker_threads) == 4
    assert len(set(worker_threads)) == 1
    assert worker_threads[0] != threading.get_ident()


def test_dense_cad_serializes_tesseract_sheet_builds_across_pages(monkeypatch):
    content = make_pdf(["CAD 1", "CAD 2", "CAD 3", "CAD 4"])
    pages = [
        ParsedPdfPage(
            page_number=page_number,
            text="",
            segments=[],
            page_type="vector",
            profile={"drawing_count": 50_000, "visual_required": True},
        )
        for page_number in range(1, 5)
    ]
    active_builds = 0
    maximum_active_builds = 0
    build_lock = threading.Lock()

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter(pages))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )

    def fake_prepare(*_args, **_kwargs):
        nonlocal active_builds, maximum_active_builds
        with build_lock:
            active_builds += 1
            maximum_active_builds = max(maximum_active_builds, active_builds)
        time.sleep(0.03)
        with build_lock:
            active_builds -= 1
        return []

    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        fake_prepare,
    )

    with pytest.raises(RuntimeError, match="没有可翻译的文字"):
        asyncio.run(
            main_module._translate_pdf_document_pipeline(
                content, "cad.pdf", 4, "th", "zh", ""
            )
        )

    assert maximum_active_builds == 1


def test_dense_cad_pipeline_writes_detected_diagonal_watermark(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "อาคาร",
        "source_confidence": 99.0,
    }
    watermark = {
        "segment_id": "pdf:p1:w1",
        "page_number": 1,
        "text": "ดําเนินการก่อสร้างแล้วเสร็จ",
        "source_kind": "outline-text",
        "bbox": (10.0, 10.0, 220.0, 160.0),
        "font_size": 18.0,
        "color": "#b80000",
        "alignment": "center",
        "metadata": {
            "diagonal_watermark": True,
            "marked_content_tag": "oc1",
            "watermark_lines": [
                [[10.0, 10.0], [220.0, 160.0]],
                [[10.0, 20.0], [220.0, 170.0]],
            ],
            "baseline_start": [10.0, 15.0],
            "baseline_end": [220.0, 165.0],
        },
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {
                "content": b"png",
                "entries": {"ID001": candidate},
                "supplemental_units": [watermark],
            }
        ],
    )
    monkeypatch.setattr(
        main_module, "detect_dense_cad_paddle_candidates", lambda *_args, **_kwargs: []
    )

    async def fake_read_indexed(*_args, **_kwargs):
        return [{"id": "ID001", "source_text": "อาคาร"}], "vision"

    async def fake_translate(segments, *_args, **_kwargs):
        values = {segment.segment_id: "建筑" for segment in segments}
        values["pdf:p1:w1"] = "已竣工"
        return values, ["openai:responses"], []

    monkeypatch.setattr(
        main_module.translator, "read_indexed_image_lines", fake_read_indexed
    )
    monkeypatch.setattr(main_module, "_translate_document_segments", fake_translate)

    _source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    written_watermark = next(
        item
        for item in result.layout_segments
        if item["segment_id"] == "pdf:p1:w1"
    )
    assert written_watermark["translated_text"] == "已竣工"
    assert written_watermark["metadata"]["diagonal_watermark"] is True


def test_pdf_pipeline_adds_uncovered_paddle_cad_text_to_translation(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    initial = {
        "bbox": (40.0, 60.0, 150.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "อาคาร",
        "source_confidence": 90.0,
    }
    supplement = {
        "bbox": (40.0, 100.0, 160.0, 122.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "ผ้าเพดาน",
        "source_confidence": 92.0,
        "candidate_provider": "paddle-full-page",
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "auto")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"initial", "entries": {"ID001": initial}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [supplement],
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_review_sheets",
        lambda _page, candidates, **_kwargs: [
            {"content": b"paddle", "entries": {"PID001": candidates[0]}}
        ],
    )

    async def fake_read_indexed(content, *_args, **_kwargs):
        if content == b"initial":
            return (
                [
                    {
                        "id": "ID001",
                        "source_text": "อาคาร",
                    }
                ],
                "vision",
            )
        assert content == b"paddle"
        return (
            [
                {
                    "id": "PID001",
                    "source_text": "ผ้าเพดาน",
                }
            ],
            "vision-paddle",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_read_indexed,
    )

    async def fake_translate_segments(segments, *_args, **_kwargs):
        translated = {
            segment.segment_id: {
                "อาคาร": "建筑",
                "ผ้าเพดาน": "吊顶",
            }[segment.text]
            for segment in segments
        }
        return translated, ["text"], []

    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        fake_translate_segments,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert "ผ้าเพดาน" in source_text
    recovered = next(
        item
        for item in result.layout_segments
        if item["text"] == "ผ้าเพดาน"
    )
    assert recovered["translated_text"] == "吊顶"
    assert recovered["metadata"]["ocr_provider"] == "paddle-full-page"


async def _fake_two_stage_cad_text_translation(segments, *_args, **_kwargs):
    return (
        {
            segment.segment_id: segment.text.replace("อาคาร", "建筑")
            for segment in segments
        },
        ["text"],
        [],
    )


def test_pdf_pipeline_reviews_only_unresolved_cad_candidates_at_high_resolution(
    monkeypatch,
):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "อาคาร",
        "source_confidence": 90.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"base", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "subset_indexed_translation_sheet",
        lambda sheet, _ids: sheet,
    )
    monkeypatch.setattr(
        main_module,
        "recognize_indexed_sheet_rows",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        main_module,
        "recover_dense_cad_review_text_lines",
        lambda *_args, **_kwargs: [],
    )

    def fake_prepare_review(_page, candidates, **kwargs):
        assert kwargs["desired_width"] == 6400
        assert len(candidates) == 1
        assert candidates[0]["origin_item_id"] == "ID001"
        return [
            {
                "content": b"review",
                "entries": {"RID0001": candidates[0]},
                "focused_review": True,
            }
        ]

    monkeypatch.setattr(
        main_module, "prepare_dense_cad_review_sheets", fake_prepare_review
    )
    calls = []

    async def fake_translate_indexed(content, *_args, **_kwargs):
        calls.append(content)
        if content != b"review":
            # A model-level no-text answer is syntactically complete, but it
            # contradicts the reliable Thai locator hint and must be retried.
            return [{"id": "ID001", "source_text": "[NO_TEXT]"}], "vision"
        return (
            [
                {
                    "id": "RID0001",
                    "source_text": "ABC อาคาร 123",
                    "translated_text": "ABC 建筑 123",
                }
            ],
            "vision-review",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert calls == [b"base", b"review"]
    assert source_text == "ABC อาคาร 123"
    assert result.layout_segments[0]["bbox"] == candidate["bbox"]
    assert result.layout_segments[0]["translated_text"] == "ABC 建筑 123"


def test_pdf_pipeline_retries_timed_out_cad_high_resolution_review(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "",
        "source_confidence": 0.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"base", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "recognize_indexed_sheet_rows",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        main_module,
        "recover_dense_cad_review_text_lines",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_review_sheets",
        lambda _page, candidates, **_kwargs: [
            {"content": b"review", "entries": {"RID0001": candidates[0]}}
        ],
    )
    calls = []

    async def fake_translate_indexed(content, *_args, **_kwargs):
        calls.append(content)
        if content == b"base":
            return [], "vision"
        if calls.count(b"review") == 1:
            raise asyncio.TimeoutError()
        return (
            [
                {
                    "id": "RID0001",
                    "source_text": "อาคาร",
                    "translated_text": "建筑",
                }
            ],
            "vision-review",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert calls == [b"base", b"review", b"review"]
    assert source_text == "อาคาร"
    assert result.layout_segments[0]["translated_text"] == "建筑"


def test_cad_duplicate_fragment_is_covered_by_confirmed_translation():
    assert main_module._cad_candidate_is_covered_by_translation(
        (105, 105, 125, 120),
        [fitz.Rect(100, 100, 180, 140)],
    )
    assert main_module._cad_candidate_is_covered_by_translation(
        (125, 105, 140, 135),
        [fitz.Rect(100, 100, 160, 120)],
    )
    assert not main_module._cad_candidate_is_covered_by_translation(
        (170, 105, 205, 120),
        [fitz.Rect(100, 100, 180, 140)],
    )


def test_pdf_pipeline_fails_when_cad_review_still_omits_candidate_ids(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "",
        "source_confidence": 0.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"base", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "recognize_indexed_sheet_rows",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_review_sheets",
        lambda _page, candidates, **_kwargs: [
            {"content": b"review", "entries": {"RID0001": candidates[0]}}
        ],
    )

    async def fake_translate_indexed(*_args, **_kwargs):
        return [], "vision"

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )

    with pytest.raises(RuntimeError, match="ID_MISMATCH.*CAD 高清复核"):
        asyncio.run(
            main_module._translate_pdf_document_pipeline(
                content, "cad.pdf", 1, "th", "zh", ""
            )
        )


def test_pdf_pipeline_preserves_only_tiny_unreadable_cad_label(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="สวัสดี",
        segments=[
            DocumentSegment(
                segment_id="pdf:p1:u1",
                page_number=1,
                text="สวัสดี",
            )
        ],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 60.0, 93.0),
        "rotation": 90,
        "vertical": True,
        "source_hint": "๐ชี 3",
        "source_confidence": 83.0,
    }

    assert main_module._cad_candidate_is_tiny_unreadable_label(candidate)
    assert main_module._cad_candidate_is_tiny_unreadable_label(
        {
            **candidate,
            "bbox": (40.0, 60.0, 56.0, 86.0),
            "source_hint": "2 ธร",
        }
    )
    assert main_module._cad_candidate_is_tiny_unreadable_label(
        {
            **candidate,
            "bbox": (40.0, 60.0, 64.0, 134.0),
            "source_hint": "ร 7",
        }
    )
    assert not main_module._cad_candidate_is_tiny_unreadable_label(
        {**candidate, "source_hint": "ประตู", "bbox": (40.0, 60.0, 120.0, 93.0)}
    )
    assert not main_module._cad_candidate_is_tiny_unreadable_label(
        {**candidate, "source_hint": "2 ผนัง"}
    )

    monkeypatch.setenv("APP_CAD_OCR_MODE", "indexed")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module.database,
        "find_exact_knowledge_many",
        lambda *_args, **_kwargs: {},
    )

    async def fake_translate_batch(batch, *_args, **_kwargs):
        return {segment.segment_id: "你好" for segment in batch}, ["text"], []

    monkeypatch.setattr(
        main_module,
        "_translate_document_segment_batch",
        fake_translate_batch,
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"base", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_review_sheets",
        lambda *_args, **_kwargs: pytest.fail("tiny label started high-resolution review"),
    )

    async def fake_translate_indexed(*_args, **_kwargs):
        return [], "vision"

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )

    _source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert [item["translated_text"] for item in result.layout_segments] == ["你好"]
    assert any("保留 1 个多工具未确认的极小 CAD 标签原文" in warning for warning in result.warnings)


def test_pdf_pipeline_fast_mode_uses_local_cad_transcription(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    unit = {
        "segment_id": "pdf:p1:o1",
        "page_number": 1,
        "text": "ABC อาคาร 123",
        "source_kind": "outline-text",
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "font_size": 12.0,
        "color": "#000000",
        "alignment": "left",
        "metadata": {
            "visual_pdf_version": 1,
            "ocr_provider": "paddleocr",
            "rotation": 0,
            "line_count": 1,
            "leading": 12.0,
            "cover_bbox": [40.0, 60.0, 210.0, 82.0],
        },
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "fast")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: pytest.fail("fast mode used indexed GPT OCR"),
    )

    def fake_extract_visual(*_args, **kwargs):
        kwargs["diagnostics"].update(seed_count=1, accepted_count=1)
        return [unit]

    monkeypatch.setattr(
        main_module,
        "extract_visual_page_units",
        fake_extract_visual,
    )

    async def fake_translate_segments(*_args, **_kwargs):
        return {"pdf:p1:o1": "ABC 建筑 123"}, ["vision"], []

    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        fake_translate_segments,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert "ABC อาคาร 123" in source_text
    assert result.layout_segments[0]["translated_text"] == "ABC 建筑 123"
    assert result.layout_segments[0]["metadata"]["ocr_provider"] == "paddleocr"


def test_pdf_pipeline_auto_mode_uses_tesseract_primary_recall_with_paddle_audit(
    monkeypatch,
):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "auto")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())

    def fake_prepare(*_args, **kwargs):
        assert kwargs["native_units"] == []
        assert kwargs["detection_provider"] == "tesseract"
        assert kwargs["rows_per_sheet"] == main_module.CAD_PADDLE_DETECTOR_ROWS_PER_SHEET
        return [{"content": b"png", "entries": {"ID001": candidate}}]

    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        fake_prepare,
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )

    async def fake_translate_indexed(*_args, **_kwargs):
        return (
            [
                {
                    "id": "ID001",
                    "source_text": "ABC อาคาร 123",
                    "translated_text": "ABC 建筑 123",
                }
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert source_text == "ABC อาคาร 123"
    assert result.layout_segments[0]["translated_text"] == "ABC 建筑 123"
    assert result.layout_segments[0]["metadata"]["ocr_provider"] == "gpt-indexed-source"
    assert result.layout_segments[0]["font_size"] <= 16.0


def test_pdf_pipeline_native_cad_uses_indexed_outline_supplement_not_deep_ocr(
    monkeypatch,
):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={
            "drawing_count": 50_000,
            "visual_required": True,
            "native_text_complete": True,
        },
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "",
        "source_confidence": 0.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "auto")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "extract_visual_page_units_in_worker",
        lambda *_args, **_kwargs: pytest.fail("native CAD started deep OCR"),
    )
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **kwargs: (
            kwargs["detection_provider"] == "tesseract"
            and [{"content": b"png", "entries": {"ID001": candidate}}]
        ),
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )

    async def fake_translate_indexed(*_args, **_kwargs):
        return (
            [
                {
                    "id": "ID001",
                    "source_text": "ABC อาคาร 123",
                    "translated_text": "ABC 建筑 123",
                }
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        _fake_two_stage_cad_text_translation,
    )

    _source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert result.layout_segments[0]["translated_text"] == "ABC 建筑 123"


def test_pdf_pipeline_reliable_tesseract_line_uses_two_stage_translation(monkeypatch):
    content = make_pdf(["CAD"])
    page = ParsedPdfPage(
        page_number=1,
        text="",
        segments=[],
        page_type="vector",
        profile={"drawing_count": 50_000, "visual_required": True},
    )
    candidate = {
        "bbox": (40.0, 60.0, 210.0, 82.0),
        "rotation": 0,
        "vertical": False,
        "source_hint": "ABC อาคาร 123",
        "source_confidence": 98.0,
    }

    monkeypatch.setenv("APP_CAD_OCR_MODE", "auto")
    monkeypatch.setattr(main_module, "iter_pdf_pages", lambda _: iter([page]))
    monkeypatch.setattr(main_module.translator, "provider", object())
    monkeypatch.setattr(
        main_module,
        "prepare_dense_cad_translation_sheets",
        lambda *_args, **_kwargs: [
            {"content": b"png", "entries": {"ID001": candidate}}
        ],
    )
    monkeypatch.setattr(
        main_module,
        "detect_dense_cad_paddle_candidates",
        lambda *_args, **_kwargs: [],
    )
    translated_sources = []

    async def fake_text_translation(segments, *_args, **_kwargs):
        translated_sources.extend(segment.text for segment in segments)
        return await _fake_two_stage_cad_text_translation(segments)

    async def fake_translate_indexed(*_args, **kwargs):
        assert kwargs["expected_sources"] == {"ID001": "ABC อาคาร 123"}
        return (
            [
                {
                    "id": "ID001",
                    "source_text": "ABC อาคาร 123",
                    "translated_text": "ABC 建筑 123",
                }
            ],
            "vision",
        )

    monkeypatch.setattr(
        main_module.translator,
        "read_indexed_image_lines",
        fake_translate_indexed,
    )
    monkeypatch.setattr(
        main_module,
        "_translate_document_segments",
        fake_text_translation,
    )

    _source_text, result = asyncio.run(
        main_module._translate_pdf_document_pipeline(
            content, "cad.pdf", 1, "th", "zh", ""
        )
    )

    assert translated_sources == ["ABC อาคาร 123"]
    assert result.layout_segments[0]["translated_text"] == "ABC 建筑 123"


def test_indexed_image_translation_accepts_tesseract_confirmed_reading_once():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "ID001",
                "source_text": "ABC อาคาร 123",
                "translated_text": "ABC 建筑 123",
            }
        ],
        "vision",
    )

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["ID001"],
            "zh",
            require_complete=False,
            verify_thai_source=True,
            expected_sources={"ID001": "ABC อาคาร 123"},
        )
    )

    assert service.provider.translate_indexed_image_lines.call_count == 1
    assert items[0]["translated_text"] == "ABC 建筑 123"


def test_indexed_image_translation_accepts_single_high_resolution_thai_reading():
    service = TranslationService(replace(settings, ai_provider="demo", openai_api_key=""))
    service.provider = AsyncMock()
    service.provider.translate_indexed_image_lines.return_value = (
        [
            {
                "id": "RID0001",
                "source_text": "นางนภัส",
                "translated_text": "那帕",
            }
        ],
        "vision",
    )

    items, _ = asyncio.run(
        service.translate_indexed_image_lines(
            b"image",
            "image/png",
            ["RID0001"],
            "zh",
            verify_thai_source=True,
            accept_single_thai_reading=True,
        )
    )

    assert service.provider.translate_indexed_image_lines.call_count == 1
    assert items == [
        {
            "id": "RID0001",
            "source_text": "นางนภัส",
            "translated_text": "那帕",
        }
    ]


def test_scanned_pdf_respects_visual_page_limit(monkeypatch):
    content = make_pdf(["Page one", "Page two"], image_only=True)
    monkeypatch.setattr(
        main_module, "settings", replace(settings, pdf_ocr_max_pages=1)
    )

    with TestClient(app) as client:
        response = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("too-many-pages.pdf", content, "application/pdf")},
        )

    assert response.status_code == 422
    assert "最多处理 1 页" in response.json()["detail"]


def test_rejects_same_language_and_unsupported_file():
    with TestClient(app) as client:
        same_language = client.post(
            "/api/translate",
            json={"text": "你好", "source_language": "zh", "target_language": "zh"},
        )
        unsupported = client.post(
            "/api/translate/image",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("image.gif", b"gif", "image/gif")},
        )

    assert same_language.status_code == 422
    assert unsupported.status_code == 422


def test_rejects_damaged_and_oversized_documents():
    with TestClient(app) as client:
        damaged = client.post(
            "/api/translate/document",
            data={"source_language": "auto", "target_language": "zh"},
            files={"file": ("broken.pdf", b"not-a-pdf", "application/pdf")},
        )
        oversized = client.post(
            "/api/translate/document",
            data={"source_language": "en", "target_language": "zh"},
            files={"file": ("large.txt", b"a" * 50001, "text/plain")},
        )

    assert damaged.status_code == 422
    assert oversized.status_code == 422


def test_language_detection_for_three_languages():
    text = "中文项目使用 Python FastAPI 和 OpenAI Responses API。" * 5
    assert detect_language(text) == "zh"
    assert detect_language("Hello, this is a translation workspace.") == "en"
    assert detect_language("สวัสดี ยินดีต้อนรับ") == "th"


def test_prompt_scopes_context_to_the_current_translation():
    prompt = translation_instructions("zh", "th", "SiamLink 是产品名，保留英文")
    assert "只用于本次翻译" in prompt
    assert "SiamLink 是产品名" in prompt
    assert "不得向译文添加原文没有的信息" in prompt

    marked_prompt = translation_instructions(
        "zh", "th", "", preserve_page_markers=True
    )
    assert "【第 N 页】页码标记必须原样保留" in marked_prompt

    thai_date_prompt = structured_translation_instructions("th", "zh", "")
    assert "2026年2月25日" in thai_date_prompt


def test_openai_provider_uses_responses_api_and_none_reasoning():
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="auto",
            openai_reasoning_effort="none",
        )
    )
    provider._post = AsyncMock(
        return_value=httpx.Response(200, json={"output_text": "สวัสดี"})
    )

    translated, route = asyncio.run(
        provider.translate("你好", "zh", "th", "SiamLink 是产品名")
    )

    assert translated == "สวัสดี"
    assert route == "openai:responses"
    path, payload = provider._post.await_args.args
    assert path == "/responses"
    assert payload["reasoning"] == {"effort": "none"}
    assert "SiamLink 是产品名" in payload["instructions"]


def test_openai_structured_segments_reserve_output_room_for_all_ids():
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="responses",
        )
    )
    provider._request_compatible = AsyncMock(
        return_value=(
            '{"translations":{"cad:p1:source:001":"墙体"}}',
            "openai:responses",
        )
    )

    translated, route = asyncio.run(
        provider.translate_segments(
            {"cad:p1:source:001": "ผนัง"}, "th", "zh", ""
        )
    )

    assert translated == {"cad:p1:source:001": "墙体"}
    assert route == "openai:responses"
    responses_payload, chat_payload = provider._request_compatible.await_args.args
    assert responses_payload["max_output_tokens"] == 12_000
    assert chat_payload["max_completion_tokens"] == 12_000


def test_openai_provider_falls_back_to_chat_when_responses_is_unsupported():
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="auto",
        )
    )
    provider._post = AsyncMock(
        side_effect=[
            httpx.Response(404, text="not found"),
            httpx.Response(200, json={"choices": [{"message": {"content": "Hello"}}]}),
        ]
    )

    translated, route = asyncio.run(provider.translate("你好", "zh", "en", ""))

    assert translated == "Hello"
    assert route == "openai:chat-completions"
    assert [call.args[0] for call in provider._post.await_args_list] == [
        "/responses",
        "/chat/completions",
    ]


def test_openai_provider_retries_rate_limit_before_returning(monkeypatch):
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="responses",
            openai_max_retries=2,
            openai_retry_base_seconds=0.25,
        )
    )
    provider._post = AsyncMock(
        side_effect=[
            httpx.Response(
                429,
                headers={"Retry-After": "0.5"},
                json={"error": {"message": "Too many pending requests"}},
            ),
            httpx.Response(200, json={"output_text": "สวัสดี"}),
        ]
    )
    sleep = AsyncMock()
    monkeypatch.setattr("app.services.translator.asyncio.sleep", sleep)

    translated, route = asyncio.run(
        provider.translate("你好", "zh", "th", "")
    )

    assert translated == "สวัสดี"
    assert route == "openai:responses"
    assert provider._post.await_count == 2
    sleep.assert_awaited_once_with(0.5)


def test_openai_provider_reports_rate_limit_after_retries(monkeypatch):
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="responses",
            openai_max_retries=1,
            openai_retry_base_seconds=0.1,
        )
    )
    provider._post = AsyncMock(
        side_effect=[
            httpx.Response(429, text="Too many pending requests"),
            httpx.Response(429, text="Too many pending requests"),
        ]
    )
    monkeypatch.setattr("app.services.translator.asyncio.sleep", AsyncMock())

    try:
        asyncio.run(provider.translate("你好", "zh", "th", ""))
    except RuntimeError as exc:
        assert str(exc) == "模型网关当前排队请求过多，系统已自动重试仍未成功，请稍后再试"
    else:
        raise AssertionError("expected rate-limit RuntimeError")


def test_openai_provider_retries_transient_gateway_auth_failure(monkeypatch):
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="responses",
            openai_max_retries=2,
            openai_retry_base_seconds=0.25,
        )
    )
    provider._post = AsyncMock(
        side_effect=[
            httpx.Response(403, text="temporary upstream channel rejection"),
            httpx.Response(200, json={"output_text": "Hello"}),
        ]
    )
    sleep = AsyncMock()
    monkeypatch.setattr("app.services.translator.asyncio.sleep", sleep)

    translated, route = asyncio.run(provider.translate("你好", "zh", "en", ""))

    assert translated == "Hello"
    assert route == "openai:responses"
    assert provider._post.await_count == 2
    sleep.assert_awaited_once_with(0.25)


def test_openai_provider_retries_transient_connection_failure(monkeypatch):
    provider = OpenAIProvider(
        replace(
            settings,
            ai_provider="openai",
            openai_api_key="test-key",
            openai_api_mode="responses",
            openai_max_retries=2,
            openai_retry_base_seconds=0.25,
        )
    )
    provider._post = AsyncMock(
        side_effect=[
            RuntimeError("无法连接模型服务，请检查网络或网关状态"),
            httpx.Response(200, json={"output_text": "Hello"}),
        ]
    )
    sleep = AsyncMock()
    monkeypatch.setattr("app.services.translator.asyncio.sleep", sleep)

    translated, route = asyncio.run(provider.translate("你好", "zh", "en", ""))

    assert translated == "Hello"
    assert route == "openai:responses"
    assert provider._post.await_count == 2
    sleep.assert_awaited_once_with(0.25)
