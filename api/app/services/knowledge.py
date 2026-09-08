import unicodedata
from io import BytesIO
from typing import Dict, List, Optional, Tuple
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


KNOWLEDGE_IMPORT_MAX_ROWS = 10000
KNOWLEDGE_XLSX_MAX_BYTES = 10 * 1024 * 1024
KNOWLEDGE_XLSX_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
KNOWLEDGE_XLSX_MAX_ARCHIVE_FILES = 5000
KNOWLEDGE_XLSX_HEADER_SCAN_ROWS = 20

_HEADER_ALIASES = {
    "thai": {
        "originaltextthai",
        "originalthai",
        "thaitext",
        "thai",
        "ภาษาไทย",
    },
    "english_translated": {
        "translatedtextenglish",
        "translatedenglish",
        "englishtext",
        "english",
    },
    "english_revised": {
        "revisedtextenglish",
        "revisedenglish",
    },
    "chinese_translated": {
        "translatedtextchinese",
        "translatedchinese",
        "chinesetext",
        "chinese",
        "中文",
    },
    "chinese_revised": {
        "revisedtextchinese",
        "revisedchinese",
        "修订中文",
    },
}


def parse_knowledge_xlsx(
    content: bytes,
) -> Tuple[List[Dict[str, str]], int, int]:
    if len(content) > KNOWLEDGE_XLSX_MAX_BYTES:
        raise ValueError("知识库 Excel 文件不能超过 10 MB")
    _validate_xlsx_archive(content)

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏且格式为 XLSX") from exc

    rows: List[Dict[str, str]] = []
    total_rows = 0
    invalid_rows = 0
    recognized_sheet = False
    try:
        for worksheet in workbook.worksheets:
            header = _find_header(worksheet)
            if header is None:
                continue
            recognized_sheet = True
            header_row, columns = header
            max_data_row = header_row + KNOWLEDGE_IMPORT_MAX_ROWS + 1
            if worksheet.max_row is None:
                worksheet.calculate_dimension(force=True)
            worksheet_max_row = worksheet.max_row or header_row
            for raw_row in worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=min(worksheet_max_row, max_data_row),
                values_only=True,
            ):
                if not raw_row or not any(_cell_text(value) for value in raw_row):
                    continue
                total_rows += 1
                if total_rows > KNOWLEDGE_IMPORT_MAX_ROWS:
                    raise ValueError(
                        f"单次最多导入 {KNOWLEDGE_IMPORT_MAX_ROWS:,} 条知识库记录"
                    )

                thai_text = _value_at(raw_row, columns.get("thai"))
                english_text = _preferred_value(
                    raw_row,
                    columns.get("english_revised"),
                    columns.get("english_translated"),
                )
                chinese_text = _preferred_value(
                    raw_row,
                    columns.get("chinese_revised"),
                    columns.get("chinese_translated"),
                )
                if not _valid_knowledge_row(thai_text, chinese_text, english_text):
                    invalid_rows += 1
                    continue
                rows.append(
                    {
                        "th": thai_text,
                        "zh": chinese_text,
                        "en": english_text,
                    }
                )
    finally:
        workbook.close()

    if not recognized_sheet:
        raise ValueError(
            "Excel 表头不符合模板：需包含 Original Text（Thai），以及英语或中文译文列"
        )
    return rows, total_rows, invalid_rows


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > KNOWLEDGE_XLSX_MAX_ARCHIVE_FILES:
                raise ValueError("Excel 文件内容过于复杂，无法导入")
            if (
                sum(item.file_size for item in files)
                > KNOWLEDGE_XLSX_MAX_UNCOMPRESSED_BYTES
            ):
                raise ValueError("Excel 文件解压后内容过大，无法导入")
    except BadZipFile as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏且格式为 XLSX") from exc


def _find_header(worksheet) -> Optional[Tuple[int, Dict[str, int]]]:
    max_row = min(worksheet.max_row or 1, KNOWLEDGE_XLSX_HEADER_SCAN_ROWS)
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_row, values_only=True),
        start=1,
    ):
        columns: Dict[str, int] = {}
        for index, value in enumerate(values):
            normalized = _normalize_header(value)
            if not normalized:
                continue
            for field, aliases in _HEADER_ALIASES.items():
                if normalized in aliases and field not in columns:
                    columns[field] = index
        has_translation = any(
            field in columns
            for field in (
                "english_translated",
                "english_revised",
                "chinese_translated",
                "chinese_revised",
            )
        )
        if "thai" in columns and has_translation:
            return row_number, columns
    return None


def _normalize_header(value) -> str:
    text = unicodedata.normalize("NFKC", _cell_text(value)).casefold()
    return "".join(character for character in text if character.isalnum())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _value_at(row, index: Optional[int]) -> str:
    if index is None or index >= len(row):
        return ""
    return _cell_text(row[index])


def _preferred_value(
    row,
    revised_index: Optional[int],
    translated_index: Optional[int],
) -> str:
    return _value_at(row, revised_index) or _value_at(row, translated_index)


def _valid_knowledge_row(
    thai_text: str,
    chinese_text: str,
    english_text: str,
) -> bool:
    return (
        0 < len(thai_text) <= 2000
        and bool(chinese_text or english_text)
        and len(chinese_text) <= 5000
        and len(english_text) <= 5000
    )
